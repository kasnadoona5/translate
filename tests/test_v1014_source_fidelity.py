from __future__ import annotations

from docx import Document as WordDocument

from tarjomeh.core.pipeline import (
    _decisions_without_regressed_edits,
    _salvage_local_refinement_edits,
    _target_paragraphs_for_alignment,
    _translated_paragraph_metadata,
)
from tarjomeh.core.prompts import CRITIQUE_PROMPT
from tarjomeh.exporters.base import TranslatedDocument, TranslatedParagraph
from tarjomeh.exporters.docx_exporter import DocxExporter
from tarjomeh.quality.structure_audit import (
    TRANSLATION_STRUCTURE_MISMATCH,
    audit_structure,
)


class _AcceptedIntegrity:
    accepted = True

    def to_dict(self) -> dict[str, object]:
        return {"accepted": True, "blocking_count": 0, "findings": []}


class _AcceptingGate:
    def evaluate(self, source: str, candidate: str, **kwargs: object):
        return _AcceptedIntegrity()


def test_exact_table_lines_recover_row_identity_without_touching_prose() -> None:
    rows = [f"ردیف {index}" for index in range(1, 83)]
    recovered, boundary = _target_paragraphs_for_alignment(
        "\n".join(rows),
        expected_count=82,
        table_like=True,
    )
    assert recovered == rows
    assert boundary == "table_line_boundaries"

    prose, prose_boundary = _target_paragraphs_for_alignment(
        "سطر نخست\nسطر دوم",
        expected_count=2,
        table_like=False,
    )
    assert prose == ["سطر نخست\nسطر دوم"]
    assert prose_boundary == "paragraph_boundaries"


def test_only_degraded_empty_table_rows_receive_export_suppression() -> None:
    flagged = _translated_paragraph_metadata(
        {"is_table": True},
        translated_text="",
        degraded_alignment=True,
    )
    deliberate_blank = _translated_paragraph_metadata(
        {"is_table": False},
        translated_text="",
        degraded_alignment=True,
    )
    assert flagged["suppress_empty_target_export"] is True
    assert "suppress_empty_target_export" not in deliberate_blank


def test_docx_skips_only_flagged_target_placeholder(tmp_path) -> None:
    output = tmp_path / "placeholder.docx"
    document = TranslatedDocument(
        paragraphs=[
            TranslatedParagraph(0, "A", "آغاز"),
            TranslatedParagraph(
                1,
                "table row",
                "",
                metadata={
                    "is_table": True,
                    "alignment_placeholder": True,
                    "suppress_empty_target_export": True,
                },
            ),
            TranslatedParagraph(2, "blank", ""),
            TranslatedParagraph(3, "B", "پایان"),
        ]
    )
    DocxExporter().export(document, output, bilingual_mode="target_only")
    texts = [paragraph.text for paragraph in WordDocument(output).paragraphs]
    assert texts == ["آغاز", "", "پایان"]


def test_regressed_edit_is_rejected_while_source_faithful_count_survives() -> None:
    previous = (
        "این فصل به سه مسئله می‌پردازد. "
        "سوم، از پیوند میان قلمروگرایی و دولت‌بودگی فراتر می‌رود."
    )
    proposed = (
        "این فصل به دو مسئله می‌پردازد. "
        "سوم، میان قلمروگرایی و دولت‌بودگی تا دولت‌بودگی را بررسی می‌کند."
    )
    issue_details = [
        {
            "issue_id": "count",
            "current_persian_quote": "سه مسئله",
        },
        {
            "issue_id": "territoriality",
            "current_persian_quote": (
                "از پیوند میان قلمروگرایی و دولت‌بودگی فراتر می‌رود"
            ),
        },
    ]
    decisions = [
        {
            "issue_id": "count",
            "decision": "accepted",
            "resulting_span": "دو مسئله",
        },
        {
            "issue_id": "territoriality",
            "decision": "accepted",
            "resulting_span": (
                "میان قلمروگرایی و دولت‌بودگی تا دولت‌بودگی را بررسی می‌کند"
            ),
        },
    ]
    safe, excluded = _decisions_without_regressed_edits(
        decisions,
        [{
            "current_persian_quote": (
                "میان قلمروگرایی و دولت‌بودگی تا دولت‌بودگی را بررسی می‌کند"
            )
        }],
    )
    final, _enriched, report = _salvage_local_refinement_edits(
        source=(
            "It addresses two issues. Third, it goes beyond the link between "
            "territoriality and statehood."
        ),
        previous=previous,
        proposed=proposed,
        issue_details=issue_details,
        issue_decisions=safe,
        integrity_gate=_AcceptingGate(),  # type: ignore[arg-type]
        protected_terms=[],
        protect_inline_english=False,
        allowed_inline_originals=[],
    )
    assert excluded == ["territoriality"]
    assert report["committed_count"] == 1
    assert "دو مسئله" in final
    assert "از پیوند میان قلمروگرایی و دولت‌بودگی فراتر می‌رود" in final
    assert "تا دولت‌بودگی را" not in final


def test_explicit_source_count_is_checked_without_a_local_enumeration() -> None:
    findings = audit_structure(
        "This chapter addresses two issues.",
        "این فصل به سه مسئله می‌پردازد.",
    )
    assert len(findings) == 1
    assert findings[0].classification == TRANSLATION_STRUCTURE_MISMATCH
    assert findings[0].check_id == "announced_count_lexical_mismatch"

    assert audit_structure(
        "This chapter addresses two issues.",
        "این فصل به دو مسئله می‌پردازد.",
    ) == []


def test_critic_contract_requires_source_counts_and_integrated_appositives() -> None:
    assert "announced quantity" in CRITIQUE_PROMPT
    assert "author's inconsistency" in CRITIQUE_PROMPT
    assert "parenthetical explanations and appositives" in CRITIQUE_PROMPT
    assert "do not interrupt or duplicate the finite predicate" in CRITIQUE_PROMPT

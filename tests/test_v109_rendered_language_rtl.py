from __future__ import annotations

import zipfile
from pathlib import Path

from docx import Document

from tarjomeh.core.pipeline import (
    audit_document_final_text,
    repair_document_source_grounded_language_artifacts,
)
from tarjomeh.core.prompts import (
    ACADEMIC_REGISTER_MODIFIER,
    CRITIQUE_PROMPT,
    REFINE_PROMPT,
)
from tarjomeh.core.term_notes import _target_spans
from tarjomeh.exporters.base import TranslatedDocument, TranslatedParagraph
from tarjomeh.exporters.docx_exporter import DocxExporter
from tarjomeh.memory.manager import _clean_style_sample
from tarjomeh.quality.integrity import (
    repair_source_grounded_language_artifacts,
    tatweel_separator_artifacts,
)


def test_source_grounded_render_repairs_cover_general_persian_artifacts() -> None:
    source = (
        "Methodological individualism and Eurocentric theory are discussed. "
        "They make history, their own and that of others."
    )
    target = (
        "\u0641\u0631\u062f\u06af\u0631\u0627\u06cc\u06cc "
        "\u0631\u0648\u0634\u200c\u0634\u0646\u0627\u062e\u062a\u06cc "
        "(methodological individualism)\u200c\u0627\u06cc \u0648 "
        "(\u0627\u0631\u0648\u067e\u0627\u0645\u062d\u0648\u0631\u06cc "
        "(Eurocentric theory)) \u0628\u0631\u0631\u0633\u06cc "
        "\u0645\u06cc\u200c\u0634\u0648\u0646\u062f. "
        "\u062a\u0627\u0631\u06cc\u062e \u0640\u0640 \u062a\u0627\u0631\u06cc\u062e "
        "\u062e\u0648\u062f \u0648 \u062f\u06cc\u06af\u0631\u0627\u0646 \u0631\u0627 "
        "\u0645\u06cc\u200c\u0633\u0627\u0632\u0646\u062f."
    )

    repaired, report = repair_source_grounded_language_artifacts(source, target)

    assert (
        "\u0631\u0648\u0634\u200c\u0634\u0646\u0627\u062e\u062a\u06cc\u200c\u0627\u06cc "
        "(methodological individualism)"
    ) in repaired
    assert (
        "(\u0627\u0631\u0648\u067e\u0627\u0645\u062d\u0648\u0631\u06cc "
        "[Eurocentric theory])"
    ) in repaired
    assert "\u062a\u0627\u0631\u06cc\u062e \u062a\u0627\u0631\u06cc\u062e" not in repaired
    assert "\u0640\u0640" not in repaired
    assert {item["type"] for item in report["repairs"]} >= {
        "parenthetical_persian_suffix",
        "nested_inline_original",
        "tatweel_separator",
        "adjacent_duplicate",
    }


def test_detached_ezafe_is_joined_without_changing_content() -> None:
    repaired, report = repair_source_grounded_language_artifacts(
        "Past and present futures of the state.",
        "\u06af\u0630\u0634\u062a\u0647 \u0648 \u062d\u0627\u0644 "
        "(\u0622\u06cc\u0646\u062f\u0647\u200c\u0647\u0627) \u06cc \u062f\u0648\u0644\u062a.",
    )
    assert repaired == (
        "\u06af\u0630\u0634\u062a\u0647 \u0648 \u062d\u0627\u0644 "
        "(\u0622\u06cc\u0646\u062f\u0647\u200c\u0647\u0627)\u06cc \u062f\u0648\u0644\u062a."
    )
    assert report["repairs"][0]["type"] == "detached_ezafe"


def test_final_document_repair_is_monotonic_and_preserves_identifiers() -> None:
    document = TranslatedDocument(paragraphs=[TranslatedParagraph(
        index=4,
        source_text="They make history, their own and others (Smith 2004).",
        translated_text=(
            "\u062a\u0627\u0631\u06cc\u062e \u0640\u0640 \u062a\u0627\u0631\u06cc\u062e "
            "\u062e\u0648\u062f \u0648 \u062f\u06cc\u06af\u0631\u0627\u0646 \u0631\u0627 "
            "\u0645\u06cc\u200c\u0633\u0627\u0632\u0646\u062f (Smith 2004)."
        ),
    )])

    report = repair_document_source_grounded_language_artifacts(document)

    assert report["accepted_repair_count"] == 2
    assert report["rejected_paragraph_count"] == 0
    assert document.paragraphs[0].translated_text.endswith("(Smith 2004).")
    assert not audit_document_final_text(document)["review_required"]


def test_modifier_inflection_can_anchor_a_multiword_first_occurrence() -> None:
    spans = _target_spans(
        "\u0627\u06cc\u0646 \u0686\u0631\u062e\u0634\u06cc "
        "\u0641\u0631\u0647\u0646\u06af\u06cc \u0645\u0647\u0645 \u0628\u0648\u062f.",
        "\u0686\u0631\u062e\u0634 \u0641\u0631\u0647\u0646\u06af\u06cc",
    )
    assert spans == [
        (4, 16, "\u0686\u0631\u062e\u0634\u06cc \u0641\u0631\u0647\u0646\u06af\u06cc")
    ]


def test_style_memory_rejects_tatweel_and_accidental_repetition() -> None:
    clean = (
        "\u0627\u06cc\u0646 \u0628\u0646\u062f \u0628\u0627 \u0646\u062b\u0631\u06cc "
        "\u0631\u0648\u0634\u0646\u060c \u062f\u0642\u06cc\u0642 \u0648 "
        "\u062f\u0627\u0646\u0634\u06af\u0627\u0647\u06cc "
        "\u0627\u0633\u062a\u062f\u0644\u0627\u0644 "
        "\u0631\u0627 \u062a\u0648\u0636\u06cc\u062d \u0645\u06cc\u200c\u062f\u0647\u062f."
    )
    assert _clean_style_sample(clean)
    assert _clean_style_sample(
        clean.replace(
            "\u0631\u0648\u0634\u0646\u060c",
            "\u0631\u0648\u0634\u0646 \u0640\u0640 \u0631\u0648\u0634\u0646\u060c",
        )
    ) == ""
    assert tatweel_separator_artifacts(
        "\u062f\u06cc\u062f\u06af\u0627\u0647 \u0646\u062e\u0633\u062a "
        "\u0640\u0640 \u062f\u06cc\u062f\u06af\u0627\u0647 \u062f\u0648\u0645"
    )


def test_contents_table_has_native_rtl_semantics_and_title_first(
    tmp_path: Path,
) -> None:
    document = TranslatedDocument(paragraphs=[
        TranslatedParagraph(
            index=0,
            source_text="Contents",
            translated_text="\u0641\u0647\u0631\u0633\u062a \u0645\u0637\u0627\u0644\u0628",
            heading_level=1,
            metadata={"structure_role": "heading"},
        ),
        TranslatedParagraph(
            index=1,
            source_text="Preface viii",
            translated_text="\u067e\u06cc\u0634\u06af\u0641\u062a\u0627\u0631 viii",
            metadata={
                "structure_role": "contents_entry",
                "toc_page_label": "viii",
                "toc_entry_kind": "unnumbered",
                "toc_level": 0,
            },
        ),
    ])
    output = tmp_path / "rtl-contents.docx"
    DocxExporter().export(document, output, "target_only")

    exported = Document(output)
    assert exported.tables[0].cell(0, 0).text == "\u067e\u06cc\u0634\u06af\u0641\u062a\u0627\u0631"
    assert exported.tables[0].cell(0, 1).text == "viii"
    with zipfile.ZipFile(output) as archive:
        xml = archive.read("word/document.xml").decode("utf-8")
    assert '<w:bidiVisual w:val="1"' in xml
    assert xml.index("\u067e\u06cc\u0634\u06af\u0641\u062a\u0627\u0631") < xml.index("viii")


def test_fluency_contract_preserves_refiner_independence() -> None:
    assert "finite" in ACADEMIC_REGISTER_MODIFIER
    assert "conceptual" in ACADEMIC_REGISTER_MODIFIER
    assert "opaque modifier stacks" in CRITIQUE_PROMPT
    assert "preserve it" in REFINE_PROMPT
    assert "suggested wording is never mandatory" in REFINE_PROMPT

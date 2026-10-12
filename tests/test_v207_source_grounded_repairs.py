"""General note/cardinality and independent-edit safeguards; no network calls."""

import asyncio
import hashlib
from collections import Counter
from types import SimpleNamespace

import pytest

from tarjomeh.chunking.chunker import Chunk
from tarjomeh.core.config import TarjomehConfig
from tarjomeh.core.pipeline import (
    TranslationPipeline,
    _critique_for_event,
    _partition_unchanged_book_term_reviews,
)
from tarjomeh.core.render_identity import compare_docx_to_rendered
from tarjomeh.exporters.base import TranslatedDocument, TranslatedParagraph
from tarjomeh.exporters.docx_exporter import DocxExporter, source_superscript_spans
from tarjomeh.glossary.book_review import check_reviewed_book_terms
from tarjomeh.jobs.database import JobDatabase
from tarjomeh.quality.critique import CritiqueResult
from tarjomeh.quality.integrity import (
    PostEditIntegrityGate,
    available_note_markers,
    extract_note_markers,
    restore_source_note_markers,
    source_note_cue_inventory,
)
from tarjomeh.quality.refiner import RefinementResult

SOURCE = "Consider (Example Author)1 - the argument follows."
TARGET = "(Example Author)¹ بخوانید ۱ — استدلال ادامه دارد."
REPAIRED = "(Example Author)¹ بخوانید — استدلال ادامه دارد."


def sha(value):
    return hashlib.sha256(value.encode("utf-8")).hexdigest()


def test_duplicate_is_counted_and_uniquely_repaired_before_hashing():
    required = Counter(extract_note_markers(SOURCE))
    assert available_note_markers(TARGET, required, source=SOURCE) == {"1": 2}
    assert not PostEditIntegrityGate().evaluate(SOURCE, TARGET).accepted
    fixed, report = restore_source_note_markers(SOURCE, TARGET)
    assert fixed == REPAIRED and report["repair_count"] == 1
    assert not report["surplus"]
    assert restore_source_note_markers(SOURCE, fixed)[0] == fixed
    assert PostEditIntegrityGate().evaluate(SOURCE, fixed).accepted


@pytest.mark.parametrize("dash", ["-", "‐", "–", "—"])
def test_source_proven_duplicate_note_handles_separated_dash_variants(dash):
    target = TARGET.replace("—", dash)
    fixed, report = restore_source_note_markers(SOURCE, target)
    assert fixed == REPAIRED.replace("—", dash) and report["repair_count"] == 1


def test_hyphenated_numeric_label_is_not_a_word_dash_note():
    target = "(Example Author)¹ رده ۱-A ثبت شد."
    assert restore_source_note_markers(SOURCE, target)[0] == target


def test_scope_uses_original_anchor_not_paragraph_position_or_unrelated_numbers():
    source = SOURCE + "\n\nChapter 1 has one problem."
    target = "فصل ۱ مسئله دارد.\n\n" + TARGET
    fixed, report = restore_source_note_markers(source, target)
    assert fixed == "فصل ۱ مسئله دارد.\n\n" + REPAIRED
    assert report["repair_count"] == 1


def test_a_source_list_can_prove_its_unique_trailing_name_anchor():
    source = "Consider (Other Writer, and Example Author)1 - the argument follows."
    assert restore_source_note_markers(source, TARGET)[0] == REPAIRED


@pytest.mark.parametrize(
    "source",
    [
        SOURCE + " It has 1 member.",
        SOURCE + " It has one member.",
        SOURCE + " It is the first option.",
        SOURCE + " Equation x = 1 holds.",
        SOURCE + " See chapter 1.",
    ],
)
def test_ordinary_source_number_is_never_deleted(source):
    fixed, report = restore_source_note_markers(source, TARGET)
    assert fixed == TARGET and report["repair_count"] == 0
    assert source_note_cue_inventory(source, TARGET)["available"] == {"1": 1}


@pytest.mark.parametrize(
    "target",
    [
        "(Example Author)¹ فصل ۱ — استدلال.",
        "(Example Author)¹ (مورد ۱ — توضیح).",
        "(Example Author)¹ x = 1 — argument.",
        "(Example Author)¹ مقدار ۱.۵ — استدلال.",
    ],
)
def test_structural_math_and_decimal_numbers_are_protected(target):
    assert restore_source_note_markers(SOURCE, target)[0] == target


@pytest.mark.parametrize(
    "source,target",
    [
        (SOURCE, "نامی¹ بخوانید ۱ — استدلال."),
        (SOURCE, "(Example Author)¹ بخوانید ۱ — ادامه ۱ — استدلال."),
        (SOURCE, "(Example Author)¹ و نامی¹ است."),
    ],
)
def test_unproved_or_multiple_surplus_cues_remain_unrepaired(source, target):
    fixed, report = restore_source_note_markers(source, target)
    assert fixed == target and report["surplus"]


def test_cross_paragraph_word_dash_is_not_removed_by_equal_counts():
    target = "(Example Author)¹\n \nبخوانید ۱ — استدلال."
    assert restore_source_note_markers(SOURCE, target)[0] == target


def test_repeated_source_notes_are_not_collapsed():
    source = "(First Author)1 - first.\n\n(Second Author)1 - second."
    target = "(First Author)¹ نخست.\n\n(Second Author)¹ دوم."
    fixed, report = restore_source_note_markers(source, target)
    assert fixed == target and not report["surplus"]


def test_one_unicode_glyph_in_a_superscript_run_is_one_physical_cue():
    text = "(Example Author)¹ استدلال."
    index = text.index("¹")
    inventory = source_note_cue_inventory(SOURCE, text, styled_spans=[(index, index + 1)])
    assert inventory["available"] == {"1": 1}
    assert set(inventory["occurrences"][0]["representations"]) == {
        "unicode_superscript",
        "docx_superscript",
    }


def test_multidigit_styled_cue_split_across_runs_counts_once():
    source = "(Example Author)12 - the argument."
    text = "(Example Author) ۱۲ استدلال."
    start = text.index("۱۲")
    inventory = source_note_cue_inventory(
        source, text, styled_spans=[(start, start + 1), (start + 1, start + 2)]
    )
    assert inventory["available"] == {"12": 1}


def test_source_proven_math_exponent_is_not_an_extra_note():
    source = SOURCE + " The expression x^1 occurs."
    text = "(Example Author)¹ عبارت x¹ است."
    inventory = source_note_cue_inventory(source, text)
    assert inventory["available"] == {"1": 1} and not inventory["surplus"]
    assert restore_source_note_markers(source, text)[0] == text


def test_source_confirmed_plain_mid_sentence_cue_still_renders_and_reads_back(tmp_path):
    source = "The third claim.3"
    target = "ادعای سوم ۳ است."
    document = TranslatedDocument(
        paragraphs=[
            TranslatedParagraph(
                index=0,
                source_text=source,
                translated_text=target,
                metadata={"superscript_markers": [{"text": "3", "relative_position": 0.5}]},
            )
        ]
    )
    path = DocxExporter().export(document, tmp_path / "plain-cue.docx")
    assert compare_docx_to_rendered(document, path, "target_only")["passed"]


def test_exporter_does_not_style_another_digit_when_unicode_cue_exists():
    metadata = {"superscript_markers": [{"text": "1", "offset": SOURCE.index("1")}]}
    assert source_superscript_spans(TARGET, metadata, SOURCE) == []


@pytest.mark.parametrize("mode", ["target_only", "inline", "side_by_side"])
def test_docx_readback_rejects_duplicate_even_when_text_identity_matches(tmp_path, mode):
    document = TranslatedDocument(
        title="Example",
        paragraphs=[
            TranslatedParagraph(
                index=0,
                source_text=SOURCE,
                translated_text=TARGET,
            )
        ],
    )
    path = DocxExporter().export(document, tmp_path / "duplicate.docx", bilingual_mode=mode)
    result = compare_docx_to_rendered(document, path, mode)
    assert not result["mismatches"] and not result["passed"]
    assert result["source_note_cues"][0]["surplus"] == ["1"]
    document.paragraphs[0].translated_text = REPAIRED
    path = DocxExporter().export(document, tmp_path / "fixed.docx", bilingual_mode=mode)
    assert compare_docx_to_rendered(document, path, mode)["passed"]


def test_superscript_style_and_unicode_are_not_double_counted_in_written_docx(tmp_path):
    from docx import Document

    document = TranslatedDocument(
        title="Example",
        paragraphs=[
            TranslatedParagraph(
                index=0,
                source_text=SOURCE,
                translated_text=REPAIRED,
            )
        ],
    )
    path = DocxExporter().export(document, tmp_path / "styled.docx")
    docx = Document(path)
    for run in docx.paragraphs[0].runs:
        if "¹" in run.text:
            run.font.superscript = True
    docx.save(path)
    assert compare_docx_to_rendered(document, path, "target_only")["passed"]


def test_inherited_superscript_style_does_not_hide_a_second_cue(tmp_path):
    from docx import Document
    from docx.enum.style import WD_STYLE_TYPE

    target = "(Example Author)¹ متن ۱ است."
    document = TranslatedDocument(
        paragraphs=[
            TranslatedParagraph(
                index=0,
                source_text=SOURCE,
                translated_text=target,
            )
        ]
    )
    path = DocxExporter().export(document, tmp_path / "inherited.docx")
    docx = Document(path)
    paragraph = docx.paragraphs[0]
    for run in list(paragraph.runs):
        paragraph._p.remove(run._r)
    style = docx.styles.add_style("SourceCue", WD_STYLE_TYPE.CHARACTER)
    style.font.superscript = True
    paragraph.add_run("(Example Author)¹ متن ")
    paragraph.add_run("۱").style = style
    paragraph.add_run(" است.")
    docx.save(path)
    audit = compare_docx_to_rendered(document, path, "target_only")
    assert not audit["mismatches"] and not audit["passed"]
    assert audit["source_note_cues"][0]["surplus"] == ["1"]


@pytest.mark.parametrize("mode", ["target_only", "inline", "side_by_side"])
def test_generated_term_note_is_not_a_duplicate_source_note(tmp_path, mode):
    ref = {
        "number": 1,
        "display_number": "۱",
        "target": "نهادها",
        "original": "institutions",
        "transliteration": "اینستیتوشنز",
    }
    document = TranslatedDocument(
        paragraphs=[
            TranslatedParagraph(
                index=0,
                source_text=SOURCE,
                translated_text=REPAIRED + " نهادها.",
                metadata={"term_note_refs": [ref]},
            )
        ],
        metadata={"term_notes": [ref]},
    )
    path = DocxExporter().export(document, tmp_path / "generated-note.docx", bilingual_mode=mode)
    audit = compare_docx_to_rendered(document, path, mode)
    assert audit["passed"] and audit["source_note_cues"][0]["available"] == {"1": 1}


def term_case():
    paragraphs = [
        "The report raises at most two issues.",
        "A national institution meets another national institution.",
    ]
    source = "\n\n".join(paragraphs)
    before = "این گزارش در نهایت فقط دو مسئله را مطرح می‌کند.\n\nنهاد ملی با نهاد ملی ملاقات می‌کند."
    after = before.replace("در نهایت فقط دو", "حداکثر دو")
    metadata = {
        "paragraph_indices": [0, 1],
        "structural_roles": ["body", "body"],
        "source_paragraph_hashes": [sha(part) for part in paragraphs],
        "source_paragraph_spans": [[0, len(paragraphs[0])], [len(paragraphs[0]) + 2, len(source)]],
    }
    matches = [
        {
            "paragraph_index": 1,
            "source_paragraph_hash": sha(paragraphs[1]),
            "source": "national institution",
            "target": "نهاد ملی",
            "source_occurrence_count": 2,
        }
    ]
    return source, before, after, metadata, matches


def test_unchanged_review_is_carried_not_resolved():
    source, before, after, metadata, matches = term_case()
    compliance, review = check_reviewed_book_terms(after, source, metadata, matches)
    assert not compliance.violations and len(review) == 1
    blocked, carried = _partition_unchanged_book_term_reviews(
        source,
        before,
        after,
        metadata,
        matches,
        review,
    )
    assert blocked == [] and carried == review
    assert check_reviewed_book_terms(after, source, metadata, matches)[1] == review


@pytest.mark.parametrize(
    "change", ["paragraph", "whitespace", "alignment", "hash", "unlocated", "count"]
)
def test_uncertain_or_changed_term_review_still_blocks(change):
    source, before, after, metadata, matches = term_case()
    if change == "paragraph":
        after += " توضیح."
    elif change == "whitespace":
        after += " "
    elif change == "alignment":
        after = after.replace("\n\n", " ")
    elif change == "hash":
        matches[0]["source_paragraph_hash"] = "wrong"
    elif change == "count":
        after = after.replace("نهاد ملی با", "نهاد با")
    _, review = check_reviewed_book_terms(after, source, metadata, matches)
    if change == "unlocated":
        review = [{"reason": "uncertain_target_paragraph_alignment"}]
    blocked, carried = _partition_unchanged_book_term_reviews(
        source,
        before,
        after,
        metadata,
        matches,
        review,
    )
    assert blocked == review and not carried


def test_independent_extra_repair_reaches_existing_critic_and_old_review_survives(tmp_path):
    source, before, after, metadata, matches = term_case()
    db = JobDatabase(tmp_path / "jobs.db")
    db.create_job("control", tmp_path / "source.txt", {})
    chunk = Chunk(
        index=0, text=source, chapter_title="Analysis", section_title="", metadata=metadata
    )
    db.save_chunks("control", [chunk])
    db.claim_worker("control", "owner")
    issue = {
        "issue_id": "bound",
        "category": "accuracy",
        "severity": "minor",
        "confidence": 0.8,
        "source_segment_id": "p1:s1",
        "source_quote": "at most two issues",
        "current_persian_quote": "در نهایت فقط دو",
        "suggested_correction": "حداکثر دو",
        "rationale": "Preserve the upper bound.",
    }

    def review(details):
        return CritiqueResult(
            accuracy=9,
            fluency=9,
            terminology=9,
            register=9,
            average=9,
            issue_details=details,
            issues=["Bound"] if details else [],
            coverage_complete=True,
            coverage_checked_segment_ids=["p1:s1", "p2:s1"],
        )

    db.log_chunk_event("control", 0, "chunk_started", {})
    db.log_chunk_event(
        "control",
        0,
        "critique_completed",
        _critique_for_event(
            review([issue]),
            9,
            -1,
            candidate_text=before,
            candidate_stage="final_retained_candidate_validation",
        ),
    )
    calls = []

    async def refine(*args, **kwargs):
        calls.append("refiner")
        return RefinementResult(
            translation=after,
            decision="revised",
            rationale="Source checked.",
            issue_decisions=[
                {
                    "issue_id": "bound",
                    "decision": "accepted",
                    "original_span": "در نهایت فقط دو",
                    "resulting_span": "حداکثر دو",
                    "rationale": "Source bound.",
                }
            ],
        )

    async def critique(*args, **kwargs):
        calls.append("critic")
        return review([])

    pipeline = object.__new__(TranslationPipeline)
    pipeline.config = TarjomehConfig()
    pipeline.config.translation.extra_final_refine_attempts = 1
    pipeline.db, pipeline.worker_id = db, "owner"
    pipeline._run_async = asyncio.run
    pipeline._canonicalize_final_translation = lambda job, idx, source, text: text
    result = pipeline._extra_final_quality_repair(
        "control",
        chunk,
        before,
        critique_tool=SimpleNamespace(critique=critique),
        refiner_tool=SimpleNamespace(refine_with_decision=refine),
        integrity_gate=PostEditIntegrityGate(),
        terminology="",
        review_context="",
        enforced_entries=[],
        enforce_auto=False,
        protected_terms=[],
        protect_inline_english=False,
        allowed_inline_originals=[],
        reviewed_matches=matches,
        threshold=9,
    )
    assert result == after and calls == ["refiner", "critic"]
    assert len(check_reviewed_book_terms(result, source, metadata, matches)[1]) == 1
    report = next(
        event["payload"]
        for event in db.get_chunk_events("control", 0)
        if event["event_type"] == "extra_final_refinement_completed"
    )
    assert report["book_term_review_count"] == 1 and report["book_term_blocking_review_count"] == 0
    assert report["book_term_unchanged_reviews"]
    assert not db.get_chunk("control", 0)["translation"]

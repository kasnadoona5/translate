"""Small post-v21 corrections must preserve state and export protection."""

from copy import deepcopy

import pytest
from docx import Document
from docx.oxml.ns import qn

from tarjomeh.chunking.chunker import Chunk
from tarjomeh.core.config import TarjomehConfig
from tarjomeh.core.render_identity import compare_docx_to_rendered
from tarjomeh.exporters.base import TranslatedDocument, TranslatedParagraph
from tarjomeh.exporters.docx_exporter import DocxExporter
from tarjomeh.memory.manager import MemoryManager
from tarjomeh.quality.evaluation import inspect_chunk

FA = (
    "\u0627\u06cc\u0646 \u06cc\u06a9 \u0627\u062f\u0639\u0627\u06cc "
    "\u0645\u0634\u0631\u0648\u0637 \u0627\u0633\u062a."
)


@pytest.mark.parametrize("policy", [0, 1])
@pytest.mark.parametrize("seeded", [False, True])
@pytest.mark.parametrize("same_chunk", [False, True])
@pytest.mark.parametrize("source,target", [
    ("A claim.", ""), ("A claim.", " \n\t"), ("", FA), (" \n\t", FA),
])
def test_empty_memory_pair_is_a_nonmutating_skip(policy, seeded, same_chunk, source, target):
    manager = MemoryManager(TarjomehConfig())
    manager.quality_policy_version = policy
    if seeded:
        manager.update_after_translation(
            Chunk(0, "A conditional claim.", "Chapter", ""), FA,
            style_approved=False,
        )
    before = deepcopy(manager.to_dict())
    pending = deepcopy(manager._pending_mapping_evidence)
    report = manager.update_after_translation(
        Chunk(0 if same_chunk else 1, source, "New chapter", "",
              metadata={"style_eligible": False, "structural_roles": ["contents"]}),
        target, style_approved=False,
    )
    assert report["status"] == "skipped_empty_memory_input"
    assert report["memory_updated"] is False
    assert report["empty_fields"] == [
        name for name, value in (("source", source), ("translation", target))
        if not value.strip()
    ]
    assert manager.to_dict() == before
    assert manager._pending_mapping_evidence == pending


def test_valid_memory_pair_keeps_all_four_layers_and_existing_report():
    manager = MemoryManager(TarjomehConfig())
    manager.quality_policy_version = 1
    report = manager.update_after_translation(
        Chunk(0, "A conditional claim.", "Chapter", ""), FA, style_approved=False,
    )
    assert report["canonical_chunk_index"] == 0
    assert len(manager.long_term._pairs) == 1
    assert manager.short_term.get_entries()[0].translation == FA
    assert manager.summary_pending
    assert manager.bilingual_summary.english_summary == ""


@pytest.mark.parametrize("mode", ["target_only", "inline", "side_by_side"])
def test_heading_sizes_and_actual_identity_are_preserved(tmp_path, mode):
    paragraphs = [
        TranslatedParagraph(level - 1, f"Heading {level}", FA + " (Example)", level)
        for level in (1, 2, 3)
    ] + [TranslatedParagraph(3, "A conditional claim.", FA)]
    document = TranslatedDocument(paragraphs=paragraphs)
    canonical = deepcopy(document)
    path = tmp_path / f"{mode}.docx"
    DocxExporter().export(document, path, bilingual_mode=mode)
    assert document == canonical
    assert compare_docx_to_rendered(document, path, mode)["passed"]
    readback = Document(path)
    units = list(readback.paragraphs) + [
        p for table in readback.tables for row in table.rows for cell in row.cells
        for p in cell.paragraphs
    ]
    headings = [p for p in units if p.style.name.startswith("Heading")]
    assert len(headings) == (3 if mode == "target_only" else 6)
    for paragraph in headings:
        level = int(paragraph.style.name.split()[-1])
        assert paragraph.style.font.size.pt == {1: 16, 2: 14, 3: 12}[level]
        assert paragraph.style.element.rPr.find(qn("w:szCs")).get(qn("w:val")) == str(
            {1: 32, 2: 28, 3: 24}[level]
        )
        for run in paragraph.runs:
            assert run.font.size is None
    for paragraph in units:
        if paragraph.text == FA and not paragraph.style.name.startswith("Heading"):
            assert all(run.font.size.pt == 11 for run in paragraph.runs)
    # The formatting fix must not make changed delivered text acceptable.
    changed = next(p for p in units if p.style.name.startswith("Heading"))
    changed.runs[0].text += " unexpected"
    readback.save(path)
    assert not compare_docx_to_rendered(document, path, mode)["passed"]


@pytest.mark.parametrize("source,target,passed", [
    ("Chapter 3", "\u0641\u0635\u0644 \u0633\u0648\u0645", True),
    ("Chapter 3", "\u0641\u0635\u0644 \u0686\u0647\u0627\u0631\u0645", False),
    ("Chapter 3", "\u0641\u0635\u0644", False),
    ("In 1973, 40% changed.", "\u06f1\u06f9\u06f7\u06f3 40%", True),
    ("In 1973, 40% changed.", "\u06f1\u06f9\u06f7\u06f4 40%", False),
    ("ISBN 978-0-123456-78-9", "ISBN 978-0-123456-78-9", True),
    ("There are 1,000 items.", "1000 items", False),
])
def test_offline_numbers_use_only_existing_conservative_equivalences(source, target, passed):
    result = inspect_chunk(source, target, "completed", [])
    check = next(item for item in result["checks"] if item["id"] == "numbers_preserved")
    assert check["passed"] is passed
    assert "localized_equivalents" in check["details"]


def test_offline_numeric_reconciliation_does_not_hide_missing_notes():
    result = inspect_chunk(
        "Chapter 3 [1].", "\u0641\u0635\u0644 \u0633\u0648\u0645.", "completed", [],
    )
    notes = next(item for item in result["checks"] if item["id"] == "note_markers_preserved")
    assert not notes["passed"] and notes["severity"] == "blocking"

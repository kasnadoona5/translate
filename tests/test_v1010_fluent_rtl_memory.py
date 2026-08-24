from __future__ import annotations

import re
import zipfile
from pathlib import Path

from docx import Document

from tarjomeh.core.pipeline import (
    _chunk_memory_admission,
    _critique_requires_refinement,
    audit_translation_language,
)
from tarjomeh.core.prompts import REFINE_PROMPT, TRANSLATE_CHUNK_PROMPT
from tarjomeh.exporters.base import TranslatedDocument, TranslatedParagraph
from tarjomeh.exporters.docx_exporter import DocxExporter
from tarjomeh.glossary.manager import GlossaryManager
from tarjomeh.memory.manager import _clean_style_sample
from tarjomeh.memory.proper_nouns import (
    ProperNouns,
    is_safe_low_authority_mapping,
)
from tarjomeh.parsers.pdf_parser import _join_block_lines
from tarjomeh.quality.critique import CritiqueResult
from tarjomeh.quality.integrity import (
    repair_source_grounded_language_artifacts,
    source_unjustified_repeated_clause_artifacts,
)


def test_grounded_structural_fluency_at_point_six_requests_existing_refiner() -> None:
    critique = CritiqueResult(
        accuracy=9,
        fluency=8,
        terminology=9,
        register=9,
        average=8.75,
        issues=["[MINOR/fluency] unclear predicate attachment"],
        issue_details=[{
            "issue_id": "fluency-attachment",
            "category": "fluency",
            "severity": "minor",
            "confidence": 0.65,
            "source_quote": "State power is a mediated condensation of forces.",
            "current_persian_quote": "قدرت دولت تراکمی میانجی گری شده از نیروهاست",
            "suggested_correction": "قدرت دولت تراکم نیروهایی است که میانجی گری می شود",
            "rationale": "The predicate and modifier attachment reproduce source order.",
        }],
    )

    assert _critique_requires_refinement(critique, 9.0)

    preference = CritiqueResult(
        accuracy=9,
        fluency=8,
        terminology=9,
        register=9,
        average=8.75,
        issues=["[MINOR/fluency] stylistic preference"],
        issue_details=[{
            **critique.issue_details[0],
            "issue_id": "fluency-preference",
            "rationale": "This alternative sounds slightly more elegant.",
        }],
    )
    assert not _critique_requires_refinement(preference, 9.0)
    assert "suggested wording is never mandatory" in REFINE_PROMPT


def test_translation_prompt_requires_accuracy_before_fluency() -> None:
    assert "every source proposition and" in TRANSLATE_CHUNK_PROMPT
    assert "finite Persian clause has an identifiable" in TRANSLATE_CHUNK_PROMPT
    assert "without merging claims" in TRANSLATE_CHUNK_PROMPT
    assert "Continuity is evidence, not authority" in TRANSLATE_CHUNK_PROMPT


def test_repeated_clause_is_review_evidence_and_not_style_authority() -> None:
    source = "The approach examines institutional variation and its consequences."
    target = (
        "این رویکرد تنوع نهادی را بررسی می‌کند و بررسی می‌کند و پیامدهای آن را می‌سنجد."
    )

    findings = source_unjustified_repeated_clause_artifacts(source, target)
    report = audit_translation_language(source, target)

    assert findings
    assert findings[0]["phrase"] == "بررسی می‌کند"
    assert report["repeated_clause_count"] == 1
    assert report["review_required"]
    assert _clean_style_sample(target) == ""


class _MemoryAdmissionDB:
    def get_chunk_events(self, job_id: str, chunk_index: int):
        return [
            {"event_type": "chunk_started", "payload": {}},
            {
                "event_type": "critique_completed",
                "payload": {
                    "valid": True,
                    "blocking_issue_count": 0,
                    "scores": {
                        "average": 8.8,
                        "accuracy": 9,
                        "fluency": 8.2,
                        "terminology": 9,
                        "register": 9,
                    },
                },
            },
        ]

    def get_chunk_review_reasons(self, job_id: str, chunk_index: int):
        return []


def test_memory_uses_configured_threshold_but_retains_continuity() -> None:
    policy = _chunk_memory_admission(
        _MemoryAdmissionDB(), "job", 0, critique_threshold=9.0
    )

    assert not policy["long_term_reliable"]
    assert policy["short_term_trust"] == "advisory_review"
    assert policy["continuity_retained"]
    assert "final_critique_below_configured_threshold" in policy["reliability_reasons"]


def test_contextual_auto_terms_are_deferred_and_loanwords_are_reclassified() -> None:
    assert not is_safe_low_authority_mapping(
        "if not on both", "و چه‌بسا بر اثر هر دو", "term"
    )
    assert not is_safe_low_authority_mapping(
        "modes of domination", "کدام شیوه‌های سلطه", "term"
    )
    assert is_safe_low_authority_mapping(
        "institutional order", "نظم نهادی", "term"
    )

    memory = ProperNouns()
    result = memory.add_noun(
        "statehood",
        "دولت‌مندی",
        category="technical_loanword",
        provenance="auto_extraction",
    )
    assert result["action"] != "ignored"
    assert memory.category_for("statehood") == "term"
    assert not memory.is_inline_eligible("statehood")


def test_legacy_contextual_mapping_is_preserved_but_prompt_deferred() -> None:
    memory = ProperNouns()
    memory.deserialize({
        "nouns": {"if not on both": "و چه‌بسا بر اثر هر دو"},
        "categories": {"if not on both": "term"},
        "provenance": {
            "if not on both": {
                "origin": "auto_extraction",
                "authority": 10,
                "observations": 1,
                "superseded": [],
            }
        },
        "introduced": [],
    })

    assert memory.all_nouns()["if not on both"] == "و چه‌بسا بر اثر هر دو"
    assert memory.is_context_deferred("if not on both")
    assert "if not on both" not in memory.get_context()


def test_persisted_auto_glossary_is_sanitized_without_touching_curated_terms() -> None:
    glossary = GlossaryManager()
    glossary.add_term("state", "دولت")
    glossary.merge_auto_extracted({
        "if not on both": {
            "target": "و چه‌بسا بر اثر هر دو",
            "category": "term",
        },
        "institutional order": {
            "target": "نظم نهادی",
            "category": "term",
        },
        "state": {
            "target": "حکومت",
            "category": "term",
        },
    })

    assert "if not on both" not in glossary
    assert "institutional order" in glossary
    assert [entry.target for entry in glossary.entries if entry.source == "state"] == [
        "دولت"
    ]


def test_nested_inline_original_uses_brackets_inside_outer_parenthesis() -> None:
    source = "For ideational institutionalism, see Smith 2004."
    target = (
        "(درباره نهادگرایی ایده‌ای (ideational institutionalism)، ر.ک. Smith 2004)"
    )

    repaired, report = repair_source_grounded_language_artifacts(source, target)

    assert repaired == (
        "(درباره نهادگرایی ایده‌ای [ideational institutionalism]، ر.ک. Smith 2004)"
    )
    assert any(item["type"] == "nested_inline_original" for item in report["repairs"])


def test_ascii_dehyphenation_uses_only_document_local_word_family_evidence() -> None:
    evidence: list[dict[str, str]] = []
    assert _join_block_lines(
        ["The argument refo-", "cuses attention."],
        known_words={"refocus"},
        dehyphenation_evidence=evidence,
    ) == "The argument refocuses attention."
    assert evidence[0]["reason"] == "morphological_family_observed_elsewhere_in_source"
    assert _join_block_lines(
        ["a well-", "defined arrangement"], known_words={"arrangement"}
    ) == "a well-defined arrangement"


def test_contents_table_uses_full_width_unequal_rtl_grid(tmp_path: Path) -> None:
    document = TranslatedDocument(paragraphs=[
        TranslatedParagraph(
            index=0,
            source_text="Contents",
            translated_text="فهرست مطالب",
            heading_level=1,
            metadata={"structure_role": "heading"},
        ),
        TranslatedParagraph(
            index=1,
            source_text="Preface viii",
            translated_text="پیشگفتار viii",
            metadata={
                "structure_role": "contents_entry",
                "toc_page_label": "viii",
                "toc_entry_kind": "unnumbered",
                "toc_level": 0,
            },
        ),
    ])
    output = tmp_path / "rtl-wide-contents.docx"
    DocxExporter().export(document, output, "target_only")

    exported = Document(output)
    table = exported.tables[0]
    with zipfile.ZipFile(output) as archive:
        xml = archive.read("word/document.xml").decode("utf-8")
    table_xml = re.search(r"<w:tbl>.*?</w:tbl>", xml).group()
    widths = [int(value) for value in re.findall(r"<w:gridCol w:w=\"(\d+)\"", table_xml)]
    table_width = int(re.search(r"<w:tblW w:type=\"dxa\" w:w=\"(\d+)\"", table_xml).group(1))

    assert table.alignment is not None
    assert '<w:bidiVisual w:val="1"' in table_xml
    assert '<w:jc w:val="right"' in table_xml
    assert '<w:tblLayout w:type="fixed"' in table_xml
    assert len(widths) == 2
    assert widths[0] > widths[1] * 4
    assert sum(widths) == table_width
    assert table_width >= 9000
    assert table.cell(0, 0).text == "پیشگفتار"
    assert table.cell(0, 1).text == "viii"

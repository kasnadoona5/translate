from unittest.mock import MagicMock

import pytest

from tarjomeh.chunking.chunker import Chunk
from tarjomeh.core.config import TarjomehConfig
from tarjomeh.core.pipeline import TranslationPipeline
from tarjomeh.glossary.book_review import _sha
from tarjomeh.jobs.database import JobDatabase
from tarjomeh.memory.manager import MemoryManager
from tarjomeh.parsers.base import Document

FA = (
    "\u0627\u06cc\u0646 \u0627\u062f\u0639\u0627 \u0627\u0647\u0645\u06cc\u062a \u062f"
    "\u0627\u0631\u062f."
)
CANDIDATE = FA.replace("\u0627\u0647\u0645\u06cc\u062a", "\u0627\u0647\u0645\u06cc\u062a\u06cc")


@pytest.mark.parametrize("accepted", [False, True])
def test_new_job_retranslation_uses_prior_context_and_atomic_outer_admission(tmp_path, accepted):
    path = tmp_path / "source.txt"
    path.write_text("This claim matters.", encoding="utf-8")
    config = TarjomehConfig()
    config.glossary.path, config.glossary.paths = "", []
    config.translation.enable_critique = False
    pipeline = TranslationPipeline(config)
    original = JobDatabase(tmp_path / "jobs.db")
    pipeline.db = original
    original.create_job("manual", path, config.to_dict())
    original.update_job_status("manual", "paused")
    original.save_job_artifact("manual", "quality_policy_v2085", {"version": 1})
    source = "This claim matters."
    chunks = [
        Chunk(
            index=i,
            text=source,
            chapter_title="Chapter",
            section_title="",
            metadata={
                "paragraph_indices": [i],
                "structural_roles": ["body"],
                "chapter_position": 1,
                "source_paragraph_spans": [[0, len(source)]],
                "source_paragraph_hashes": [_sha(source)],
            },
        )
        for i in range(3)
    ]
    original.save_chunks("manual", chunks)
    memory = MemoryManager(config)
    memory.quality_policy_version = 1
    for chunk in chunks:
        memory.update_after_translation(chunk, FA, style_approved=False)
        original.commit_chunk_checkpoint("manual", chunk.index, "completed", FA, memory.to_dict())
    previous = original.get_memory_state("manual")
    before = original.get_chunk("manual", 1)
    pipeline._parse_and_chunk = MagicMock(
        return_value=(Document(title="Book", chapters=[]), chunks)
    )

    def translate(**kwargs):
        assert kwargs["translations"] == {0: FA}
        assert [pair["chunk_index"] for pair in kwargs["memory_manager"].long_term.serialize()] == [
            0
        ]
        pipeline._pause_if_requested("manual", stage="chunk_start", chunk_index=1)
        pipeline.db.update_chunk("manual", 1, "translated", "partial unadmitted")
        assert original.get_chunk("manual", 1) == before
        pipeline.db.log_chunk_event("manual", 1, "chunk_started", {})
        return CANDIDATE if accepted else ""

    pipeline._translate_single_chunk = MagicMock(side_effect=translate)
    try:
        result = pipeline.retranslate_chunk("manual", 1)
        assert pipeline.db is original
        if accepted:
            assert result == CANDIDATE
            pairs = original.get_memory_state("manual")["past_translations"]
            assert len(pairs) == 4 and sum(pair.get("active", True) for pair in pairs) == 3
        else:
            assert result == FA
            assert original.get_memory_state("manual") == previous
            assert original.get_chunk("manual", 1) == before
        assert original.get_chunk("manual", 2)["translation"] == FA
        assert original.get_worker_lease("manual")["state"] == "released"
    finally:
        pipeline.close()


@pytest.mark.parametrize("notes_mode", ["inline", "footnote", "endnote", "both"])
def test_fresh_curated_review_preserves_scoped_persian_original_through_docx(tmp_path, notes_mode):
    from docx import Document as DocxDocument

    source = tmp_path / "source.txt"
    source.write_text(
        "The civic protocol organizes discussion and coordinates the participants. "
        "Every participant has a clear responsibility, and this responsibility "
        "remains constant throughout the discussion.",
        encoding="utf-8",
    )
    glossary = tmp_path / "terms.csv"
    term = "\u0631\u0648\u06cc\u0647"
    glossary.write_text(
        "source,target,context,sense\ncivic protocol,"
        + term
        + ",absent descriptive cue,unrelated-sense\n",
        encoding="utf-8",
    )
    target = (
        "\u0627\u06cc\u0646 " + term + " (civic protocol) "
        "\u06af\u0641\u062a\u200c\u0648\u06af\u0648 \u0631\u0627 \u0633\u0627\u0645"
        "\u0627\u0646 \u0645\u06cc\u200c\u062f\u0647\u062f \u0648 "
        "\u0645\u0634\u0627\u0631\u06a9\u062a\u200c\u06a9\u0646\u0646\u062f\u06af"
        "\u0627\u0646 \u0631\u0627 \u0647\u0645\u0627\u0647\u0646\u06af \u0645\u06cc"
        "\u200c\u06a9\u0646\u062f. "
        "\u0647\u0631 \u0645\u0634\u0627\u0631\u06a9\u062a\u200c\u06a9\u0646\u0646"
        "\u062f\u0647 \u0645\u0633\u0626\u0648\u0644\u06cc\u062a \u0645\u0634\u062e"
        "\u0635\u06cc \u062f\u0627\u0631\u062f "
        "\u0648 \u0627\u06cc\u0646 \u0645\u0633\u0626\u0648\u0644\u06cc\u062a \u062f"
        "\u0631 \u0633\u0631\u0627\u0633\u0631 \u0628\u062d\u062b \u062b\u0627\u0628"
        "\u062a \u0645\u06cc\u200c\u0645\u0627\u0646\u062f."
    )
    config = TarjomehConfig()
    config.glossary.path, config.glossary.paths = str(glossary), []
    config.glossary.enable_auto_extraction = False
    config.translation.review_book_terms_before_translating = True
    config.translation.enable_book_research = False
    config.translation.enable_web_context = False
    config.translation.enable_critique = False
    config.translation.enable_back_translation = False
    config.memory.enable_4layer = False
    config.output.format, config.output.term_notes = "docx", notes_mode
    pipeline = TranslationPipeline(config)
    pipeline.db = JobDatabase(tmp_path / "jobs.db")
    calls = []

    def completion(*args, **kwargs):
        calls.append(kwargs.get("_operation"))
        return '{"terms": []}' if kwargs.get("_operation") == "auto_term_extraction" else target

    pipeline.llm_client.complete = MagicMock(side_effect=completion)
    output = tmp_path / "translated.docx"
    try:
        pipeline.run(source, output, job_id="curated")
        assert not output.exists()
        review = pipeline.db.get_job_artifact("curated", "book_term_review_v1")
        assert review["review_version"] == 4
        proposal = next(item for item in review["proposals"] if item["source"] == "civic protocol")
        assert proposal["target"] == term and proposal["origin"] == "curated_glossary"
        assert proposal["scope_mode"] == "book_wide_recommended"
        assert len(proposal["sense_id"]) <= 60
        proposal.update(status="approved", keep_original=True)
        review["phase"] = "started"
        assert pipeline.db.complete_book_term_review("curated", review)
        pipeline.run(source, output, job_id="curated")
        text = "\n".join(paragraph.text for paragraph in DocxDocument(output).paragraphs)
        assert text.count("(civic protocol)") == 1 and term + " (civic protocol)" in text
        events = pipeline.db.get_chunk_events("curated", 0)
        final = next(
            event["payload"]
            for event in reversed(events)
            if event["event_type"] == "book_term_final_compliance"
        )
        assert final["compliant"] and not final["review"]
        assert (
            "civic protocol" not in pipeline.db.get_memory_state("curated")["proper_nouns"]["nouns"]
        )
        assert calls.count("auto_term_extraction") == 1
    finally:
        pipeline.close()

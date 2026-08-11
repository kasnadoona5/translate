from __future__ import annotations

import tempfile
import unittest
from pathlib import Path
from unittest.mock import MagicMock, patch

from tarjomeh.chunking.chunker import Chunk, SemanticChunker
from tarjomeh.context.search_providers import SearchResult, rank_search_results
from tarjomeh.core.config import TarjomehConfig
from tarjomeh.core.paragraph_protocol import decode_paragraphs, encode_paragraphs
from tarjomeh.core.pipeline import _align_chunk_translation
from tarjomeh.exporters.base import TranslatedDocument, TranslatedParagraph
from tarjomeh.exporters.docx_exporter import DocxExporter
from tarjomeh.memory.bilingual_summary import BilingualSummary
from tarjomeh.memory.manager import MemoryManager
from tarjomeh.parsers.base import Chapter, Document, Paragraph, Section
from tarjomeh.parsers.pdf_parser import PyMuPDFParser
from tarjomeh.parsers.pdf_parser import _join_block_lines
from tarjomeh.parsers.pdf_parser import _prepare_document_blocks
from tarjomeh.parsers.pdf_parser import _split_indented_paragraph_lines
from tarjomeh.parsers.pdf_parser import _sort_page_blocks_reading_order


def _pdf_block(text: str, y0: float, y1: float, *, size: float = 10.0) -> dict:
    return {
        "type": 0,
        "bbox": [50, y0, 550, y1],
        "lines": [{
            "bbox": [50, y0, 550, y1],
            "spans": [{"text": text, "size": size, "font": "Times", "flags": 0}],
        }],
    }


class TestPdfStructureReconstruction(unittest.TestCase):
    def test_ascii_hyphenated_compounds_are_not_silently_rewritten(self) -> None:
        self.assertEqual(
            _join_block_lines(["a well-", "defined arrangement"]),
            "a well-defined arrangement",
        )

    def test_two_column_blocks_read_down_each_column(self) -> None:
        def block(text: str, bbox: tuple[float, float, float, float]) -> dict:
            return {"text": text, "bbox": bbox, "orientation": "horizontal"}

        blocks = [
            block("right two", (320, 200, 500, 240)),
            block("left one", (40, 100, 220, 140)),
            block("right one", (320, 100, 500, 140)),
            block("left two", (40, 200, 220, 240)),
        ]
        ordered = _sort_page_blocks_reading_order(blocks, 560)
        self.assertEqual(
            [item["text"] for item in ordered],
            ["left one", "left two", "right one", "right two"],
        )

    def test_indented_print_paragraphs_are_recovered_conservatively(self) -> None:
        def line(text: str, x0: float, y0: float) -> dict:
            return {
                "text": text,
                "bbox": (x0, y0, 370.0, y0 + 12.0),
                "direction": (1.0, 0.0),
                "font_size": 10.0,
                "span_count": 1,
                "large_gap_count": 0,
            }

        lines = [
            line("Third, institutional approaches begin here and", 67.0, 100.0),
            line("continue at the dominant margin until they end.", 55.0, 112.0),
            line("Fourth, agent-centred approaches begin here", 67.0, 124.0),
            line("and continue without losing their boundary.", 55.0, 136.0),
        ]

        groups = _split_indented_paragraph_lines(lines)

        self.assertEqual(len(groups), 2)
        self.assertTrue(groups[0][0]["text"].startswith("Third,"))
        self.assertTrue(groups[1][0]["text"].startswith("Fourth,"))

    def test_indentation_without_sentence_boundary_is_not_split(self) -> None:
        lines = [
            {
                "text": "A sentence continues through",
                "bbox": (55.0, 100.0, 370.0, 112.0),
                "direction": (1.0, 0.0),
                "font_size": 10.0,
            },
            {
                "text": "Capital despite an incidental indent",
                "bbox": (67.0, 112.0, 370.0, 124.0),
                "direction": (1.0, 0.0),
                "font_size": 10.0,
            },
            {
                "text": "and reaches its actual ending.",
                "bbox": (55.0, 124.0, 370.0, 136.0),
                "direction": (1.0, 0.0),
                "font_size": 10.0,
            },
        ]

        self.assertEqual(len(_split_indented_paragraph_lines(lines)), 1)

    def test_structure_v3_splits_without_changing_v2_resume_behavior(self) -> None:
        page = MagicMock()
        page.rect.height = 800
        page.rect.width = 600
        page.get_text.return_value = {"blocks": [{
            "type": 0,
            "bbox": (55.0, 100.0, 370.0, 148.0),
            "lines": [
                {
                    "bbox": (67.0, 100.0, 370.0, 112.0),
                    "dir": (1.0, 0.0),
                    "spans": [{
                        "text": "Third, institutional approaches begin here and",
                        "size": 10.0,
                        "font": "Times",
                        "flags": 0,
                        "bbox": (67.0, 100.0, 370.0, 112.0),
                    }],
                },
                {
                    "bbox": (55.0, 112.0, 370.0, 124.0),
                    "dir": (1.0, 0.0),
                    "spans": [{
                        "text": "continue at the dominant margin until they end.",
                        "size": 10.0,
                        "font": "Times",
                        "flags": 0,
                        "bbox": (55.0, 112.0, 370.0, 124.0),
                    }],
                },
                {
                    "bbox": (67.0, 124.0, 370.0, 136.0),
                    "dir": (1.0, 0.0),
                    "spans": [{
                        "text": "Fourth, agent-centred approaches begin here",
                        "size": 10.0,
                        "font": "Times",
                        "flags": 0,
                        "bbox": (67.0, 124.0, 370.0, 136.0),
                    }],
                },
                {
                    "bbox": (55.0, 136.0, 370.0, 148.0),
                    "dir": (1.0, 0.0),
                    "spans": [{
                        "text": "and continue without losing their boundary.",
                        "size": 10.0,
                        "font": "Times",
                        "flags": 0,
                        "bbox": (55.0, 136.0, 370.0, 148.0),
                    }],
                },
            ],
        }]}
        document = MagicMock()
        document.page_count = 1
        document.__getitem__.return_value = page

        legacy_pages, legacy_audit = _prepare_document_blocks(
            document, structure_version=2
        )
        current_pages, current_audit = _prepare_document_blocks(
            document, structure_version=3
        )

        self.assertEqual(len(legacy_pages[0]), 1)
        self.assertEqual(legacy_audit["internal_paragraph_split_count"], 0)
        self.assertEqual(len(current_pages[0]), 2)
        self.assertEqual(current_audit["internal_paragraph_split_count"], 1)
        self.assertEqual(
            [block["source_fragment_id"] for block in current_pages[0]],
            ["pg0001.b0000.p01", "pg0001.b0000.p02"],
        )

    @patch("tarjomeh.parsers.pdf_parser.fitz.open")
    def test_chapter_heading_is_ordered_once_before_cross_page_body(
        self, mock_open: MagicMock
    ) -> None:
        document = MagicMock()
        mock_open.return_value = document
        document.name = "book.pdf"
        document.metadata = {"title": "Book", "author": "Author"}
        document.page_count = 2
        document.get_toc.return_value = [[1, "Preface", 1]]

        page_one = MagicMock()
        page_one.rect.height = 800
        page_one.rect.width = 600
        page_one.get_images.return_value = []
        page_one.get_text.side_effect = lambda kind, **kwargs: (
            "body text" if kind == "text" else {"blocks": [
                _pdf_block("The argument continues through", 600, 750),
                _pdf_block("Preface", 90, 120, size=18),
            ]}
        )
        page_two = MagicMock()
        page_two.rect.height = 800
        page_two.rect.width = 600
        page_two.get_images.return_value = []
        page_two.get_text.side_effect = lambda kind, **kwargs: (
            "body text" if kind == "text" else {"blocks": [
                _pdf_block("Preface ix", 15, 30),
                _pdf_block("several institutional settings.", 55, 100),
            ]}
        )
        document.__getitem__.side_effect = lambda index: [page_one, page_two][index]

        with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as handle:
            path = Path(handle.name)
        try:
            parsed = PyMuPDFParser().parse(path)
        finally:
            path.unlink()

        self.assertEqual(parsed.all_paragraphs[0].text, "Preface")
        self.assertEqual(parsed.all_paragraphs[0].heading_level, 1)
        self.assertEqual(
            parsed.all_paragraphs[1].text,
            "The argument continues through several institutional settings.",
        )
        self.assertEqual(
            sum(paragraph.text == "Preface" for paragraph in parsed.all_paragraphs),
            1,
        )
        self.assertTrue(parsed.all_paragraphs[1].metadata["cross_page_join"])
        self.assertTrue(parsed.all_paragraphs[0].metadata["paragraph_id"])

    @patch("tarjomeh.parsers.pdf_parser.fitz.open")
    def test_recurrent_furniture_removed_and_cross_page_prose_joined(
        self, mock_open: MagicMock
    ) -> None:
        document = MagicMock()
        mock_open.return_value = document
        document.name = "book.pdf"
        document.metadata = {"title": "Book", "author": "Author"}
        document.page_count = 4
        document.get_toc.return_value = [[1, "Introduction", 1]]

        pages = []
        bodies = [
            ("The argument offers various", 690, 735),
            ("kinds of institutional explanation.", 55, 100),
            ("A unique section title", 70, 100),
            ("Further body prose concludes the discussion.", 90, 150),
        ]
        for index, (body, y0, y1) in enumerate(bodies, 1):
            page = MagicMock()
            page.rect.height = 800
            page.get_images.return_value = []
            page.find_tables.return_value.tables = []
            blocks = [
                _pdf_block(f"Introduction {index}", 15, 30),
                _pdf_block(body, y0, y1),
            ]
            page.get_text.side_effect = lambda kind, _blocks=blocks, **kwargs: (
                "body text" if kind == "text" else {"blocks": _blocks}
            )
            pages.append(page)
        document.__getitem__.side_effect = lambda index: pages[index]

        with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as handle:
            path = Path(handle.name)
        try:
            parsed = PyMuPDFParser().parse(path)
        finally:
            path.unlink()

        texts = [paragraph.text for paragraph in parsed.all_paragraphs]
        self.assertFalse(any(text.startswith("Introduction ") for text in texts))
        self.assertIn(
            "The argument offers various kinds of institutional explanation.",
            texts,
        )
        self.assertIn("A unique section title", texts)
        audit = parsed.metadata["pdf_structure_audit"]
        self.assertEqual(audit["removed_furniture_count"], 4)
        joined = next(text for text in parsed.all_paragraphs if "various kinds" in text.text)
        self.assertTrue(joined.metadata["cross_page_join"])


class TestParagraphIdentity(unittest.TestCase):
    def test_markers_require_identity_and_order(self) -> None:
        encoded, markers = encode_paragraphs("Heading\n\nFirst body.\n\nSecond body.")
        self.assertIn("[[P0001]]", encoded)
        valid = decode_paragraphs(
            "[[P0001]]\nعنوان\n\n[[P0002]]\nبند اول.\n\n[[P0003]]\nبند دوم.",
            markers,
        )
        self.assertTrue(valid.valid)
        self.assertEqual(len(valid.paragraphs), 3)

        invalid = decode_paragraphs(
            "[[P0002]]\nبند اول.\n\n[[P0001]]\nعنوان", markers
        )
        self.assertFalse(invalid.valid)
        self.assertIn("missing_paragraph_marker", invalid.errors)

    def test_new_chunks_never_use_proportional_redistribution(self) -> None:
        paragraphs = [Paragraph("One."), Paragraph("Two.")]
        with self.assertRaises(ValueError):
            _align_chunk_translation(
                original_paragraphs=paragraphs,
                para_indices=[0, 1],
                tgt_paras=["یک بند"],
                chunk_translation="یک بند",
                strict_paragraph_identity=True,
            )


class TestContextHygiene(unittest.TestCase):
    def test_style_uses_only_qa_safe_body_prose(self) -> None:
        manager = MemoryManager(TarjomehConfig())
        front = Chunk(
            index=0,
            text="Copyright 2026. All rights reserved. ISBN 123.",
            chapter_title="Front matter",
            section_title="",
            metadata={"style_eligible": False, "structural_roles": ["front_matter"]},
        )
        policy = manager.update_after_translation(
            front, "حق نشر محفوظ است.", quality_approved=True
        )
        self.assertFalse(policy["style_sample_added"])

        prose = Chunk(
            index=1,
            text=("Institutional analysis explains historically specific relations. " * 6),
            chapter_title="Chapter 1",
            section_title="",
            metadata={"style_eligible": True, "structural_roles": ["body"]},
        )
        translation = "تحلیل نهادی مناسبات تاریخی مشخص را توضیح می‌دهد. " * 6
        policy = manager.update_after_translation(
            prose, translation, quality_approved=True
        )
        self.assertTrue(policy["style_sample_added"])
        self.assertEqual(len(manager.style_samples), 1)

    def test_reviewed_chunk_remains_advisory_short_term_context(self) -> None:
        manager = MemoryManager(TarjomehConfig())
        chunk = Chunk(
            index=0,
            text="The disputed paragraph still carries the argument forward.",
            chapter_title="Chapter 1",
            section_title="",
            metadata={
                "style_eligible": True,
                "style_body_paragraphs": [0],
                "structural_roles": ["body"],
            },
        )
        policy = manager.update_after_translation(
            chunk,
            "This translated paragraph remains useful context.",
            quality_approved=False,
        )
        context = manager.get_context_for_chunk(chunk)
        self.assertEqual(policy["short_term_trust"], "advisory_review")
        self.assertIn("needs review", context.short_term)
        self.assertIn("argument and references", context.short_term)
        self.assertEqual(
            context.references["short_term_trust"], ["advisory_review"]
        )
        self.assertFalse(policy["style_sample_added"])

    def test_style_sample_uses_only_complete_body_sentences(self) -> None:
        manager = MemoryManager(TarjomehConfig())
        chunk = Chunk(
            index=0,
            text=("Chapter title\n\n" + "Substantive body prose. " * 20),
            chapter_title="Chapter 1",
            section_title="",
            metadata={
                "style_eligible": False,
                "style_body_paragraphs": [1],
                "structural_roles": ["heading", "body"],
            },
        )
        translation = (
            "\u0639\u0646\u0648\u0627\u0646 \u062a\u0631\u062c\u0645\u0647\u200c\u0634\u062f\u0647\n\n"
            + "\u0627\u06cc\u0646 \u06cc\u06a9 \u062c\u0645\u0644\u0647 \u06a9\u0627\u0645\u0644 \u0648 \u062f\u0642\u06cc\u0642 \u062f\u0627\u0646\u0634\u06af\u0627\u0647\u06cc \u0627\u0633\u062a. " * 20
        )
        policy = manager.update_after_translation(
            chunk, translation, quality_approved=True
        )
        self.assertTrue(policy["style_sample_added"])
        self.assertNotIn("\u0639\u0646\u0648\u0627\u0646 \u062a\u0631\u062c\u0645\u0647\u200c\u0634\u062f\u0647", manager.style_samples[0])
        self.assertTrue(manager.style_samples[0].endswith("."))

    def test_html_is_removed_from_running_summary(self) -> None:
        summary = BilingualSummary()
        summary.update(
            "## English Summary\nA translated claim.\n"
            "## خلاصه فارسی\n<p dir=\"rtl\">یک گزاره ترجمه‌شده.</p>"
        )
        self.assertEqual(summary.persian_summary, "یک گزاره ترجمه‌شده.")
        self.assertNotIn("<p", summary.get_context())

    def test_search_rejects_wrong_identity(self) -> None:
        results = [
            SearchResult(
                title="Mira Valen statistics",
                url="https://example.com/football",
                snippet="The football player's latest match statistics.",
            ),
            SearchResult(
                title="Nora Valen and critical state theory",
                url="https://university.edu/paper",
                snippet="Nora Valen discusses materialist theories of the state.",
            ),
        ]
        accepted, audit = rank_search_results(
            "Nora Valen Persian academic spelling",
            results,
            identity="Nora Valen",
        )
        self.assertEqual(len(accepted), 1)
        self.assertIn("Nora", accepted[0].title)
        self.assertFalse(audit[0]["accepted"])

    def test_long_book_identity_allows_strong_partial_match(self) -> None:
        results = [SearchResult(
            title="Institutions Across Time and Space by Lena Orlov",
            url="https://www.politybooks.com/bookdetail/institutions",
            snippet="A study of institutions, history, and political organization.",
        )]
        accepted, audit = rank_search_results(
            "Lena Orlov Institutions Across Time and Space academic review",
            results,
            identity="Institutions Across Time and Space Lena Orlov Polity 2016",
        )
        self.assertEqual(accepted, results)
        self.assertTrue(audit[0]["accepted"])


class TestChapterFormatting(unittest.TestCase):
    def test_semantic_chunks_carry_structure_policy(self) -> None:
        document = Document(
            title="Book",
            chapters=[Chapter(
                title="Chapter 1",
                number=1,
                sections=[Section("", 2, [Paragraph(
                    "Body prose.",
                    metadata={"structure_role": "body", "chapter_start": True},
                )])],
            )],
        )
        chunk = SemanticChunker(max_tokens=100).chunk(document)[0]
        self.assertEqual(chunk.metadata["paragraph_protocol_version"], 1)
        self.assertTrue(chunk.metadata["style_eligible"])

    def test_docx_starts_second_chapter_on_new_page(self) -> None:
        try:
            import docx
        except ImportError:
            self.skipTest("python-docx is unavailable")
        translated = TranslatedDocument(
            title="Book",
            metadata={"chapter_page_breaks": True},
            paragraphs=[
                TranslatedParagraph(
                    0, "First.", "اول.", metadata={"chapter_position": 1, "chapter_start": True}
                ),
                TranslatedParagraph(
                    1, "Second.", "دوم.", metadata={"chapter_position": 2, "chapter_start": True}
                ),
            ],
        )
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "chapters.docx"
            DocxExporter().export(translated, path)
            rendered = docx.Document(path)
            self.assertFalse(bool(rendered.paragraphs[0].paragraph_format.page_break_before))
            self.assertTrue(bool(rendered.paragraphs[1].paragraph_format.page_break_before))


if __name__ == "__main__":
    unittest.main()

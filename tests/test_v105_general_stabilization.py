# Escaped Persian literals are intentionally kept ASCII-safe for portable tests.
# ruff: noqa: E501
from __future__ import annotations

from pathlib import Path
from tempfile import TemporaryDirectory

from docx import Document

from tarjomeh.chunking.chunker import Chunk
from tarjomeh.core import pipeline as pipeline_module
from tarjomeh.core.config import TarjomehConfig
from tarjomeh.core.pipeline import (
    _accepted_grounded_inline_originals,
    _canonicalize_chunk_inline_originals,
    _chunk_style_approved,
    _reconcile_current_entity_anchors,
)
from tarjomeh.core.term_notes import audit_inline_english_originals
from tarjomeh.exporters.base import TranslatedDocument, TranslatedParagraph
from tarjomeh.exporters.docx_exporter import DocxExporter, source_superscript_spans
from tarjomeh.memory.manager import MemoryManager
from tarjomeh.memory.proper_nouns import ProperNouns, is_usable_observed_mapping
from tarjomeh.parsers.pdf_parser import _line_superscript_markers
from tarjomeh.quality.integrity import unexpected_latin_prose
from tarjomeh.quality.structure_audit import (
    SOURCE_ANOMALY_PRESERVED,
    audit_structure,
    enumeration_length,
)


class _EventDatabase:
    def __init__(self, events: list[dict]) -> None:
        self.events = events

    def get_chunk_events(self, _job_id: str, _chunk_index: int) -> list[dict]:
        return self.events


def test_exact_unknown_entity_anchor_becomes_durable_inline_evidence() -> None:
    assert pipeline_module.has_exact_observed_anchor(
        "\u06cc\u0648\u067e \u0627\u0633\u0631 (Jupp Esser)",
        "Jupp Esser",
    )
    memory = MemoryManager(TarjomehConfig())
    source = "Jupp Esser emphasized empirical testing."
    translation = (
        "\u06cc\u0648\u067e \u0627\u0633\u0631 (Jupp Esser) \u0628\u0631 "
        "\u0622\u0632\u0645\u0648\u0646 \u062a\u062c\u0631\u0628\u06cc \u062a\u0623\u06a9\u06cc\u062f \u06a9\u0631\u062f."
    )

    result = _reconcile_current_entity_anchors(
        memory,
        ["Jupp Esser"],
        translation,
        {"Jupp Esser": "source_entity_candidate"},
        source,
    )

    assert result["observed_count"] == 1
    assert memory.proper_nouns.category_for("Jupp Esser") == "source_grounded_entity"
    assert memory.proper_nouns.is_inline_eligible("Jupp Esser")

    document = TranslatedDocument(paragraphs=[TranslatedParagraph(
        index=0,
        source_text=source,
        translated_text=translation,
    )])
    audit_inline_english_originals(
        document, memory.proper_nouns.inline_eligible_nouns()
    )
    assert "(Jupp Esser)" in document.paragraphs[0].translated_text


def test_source_grounded_technical_loanword_is_not_book_specific() -> None:
    source = "The state may be examined as a dispositif or assemblage."
    translation = (
        "\u062f\u0648\u0644\u062a \u0631\u0627 \u0645\u06cc\u200c\u062a\u0648\u0627\u0646 \u062f\u06cc\u0633\u067e\u0648\u0632\u06cc\u062a\u06cc\u0641 "
        "(dispositif) \u06cc\u0627 \u0622\u0633\u0627\u0645\u0628\u0644\u0627\u0698 "
        "(assemblage) \u062f\u0627\u0646\u0633\u062a."
    )

    categories = _accepted_grounded_inline_originals(source, translation)

    assert categories == {
        "dispositif": "technical_loanword",
        "assemblage": "technical_loanword",
    }


def test_low_authority_memory_rejects_context_fragments_but_not_curated_entries() -> None:
    unsafe = (
        (
            "United Kingdom",
            "\u062f\u0631 \u0628\u0631\u06cc\u062a\u0627\u0646\u06cc\u0627",
            "place",
        ),
        (
            "Ngai-Ling Sum",
            "\u06a9\u0645\u200c\u0627\u0647\u0645\u06cc\u062a\u200c\u062a\u0631\u060c "
            "\u0646\u06af\u0627\u06cc\u200c\u0644\u06cc\u0646\u06af \u0633\u0648\u0645",
            "person",
        ),
        (
            "synchronic",
            "\u0633\u06cc\u0646\u06a9\u0631\u0648\u0646\u06cc\u06a9\u200c\u062a\u0631",
            "term",
        ),
    )
    for source, target, category in unsafe:
        assert not is_usable_observed_mapping(source, target, category)

    memory = ProperNouns()
    assert (
        memory.add_noun(
            "United Kingdom",
            "\u062f\u0631 \u0628\u0631\u06cc\u062a\u0627\u0646\u06cc\u0627",
            category="place",
            provenance="observed_translation",
        )["action"]
        == "ignored"
    )
    assert (
        memory.add_noun(
            "United Kingdom",
            "\u0628\u0631\u06cc\u062a\u0627\u0646\u06cc\u0627",
            category="place",
            provenance="curated_glossary",
        )["action"]
        == "added"
    )


def test_resume_discards_stale_automatic_fragments_and_keeps_curated_memory() -> None:
    memory = ProperNouns()
    memory.deserialize(
        {
            "nouns": {
                "United Kingdom": "\u062f\u0631 \u0628\u0631\u06cc\u062a\u0627\u0646\u06cc\u0627",
                "state": "\u062f\u0648\u0644\u062a",
            },
            "categories": {"United Kingdom": "place", "state": "term"},
            "provenance": {
                "United Kingdom": {"origin": "observed_translation"},
                "state": {"origin": "curated_glossary"},
            },
        }
    )

    assert memory.all_nouns() == {"state": "\u062f\u0648\u0644\u062a"}


def test_canonical_chunk_text_is_shared_by_memory_and_export() -> None:
    memory = MemoryManager(TarjomehConfig())
    memory.proper_nouns.add_noun(
        "Jupp Esser",
        "\u06cc\u0648\u067e \u0627\u0633\u0631",
        category="source_grounded_entity",
        provenance="curated_glossary",
    )
    chunk = Chunk(
        0,
        "Jupp Esser emphasized empirical testing.",
        "",
        "",
    )
    initial = (
        "\u06cc\u0648\u067e \u0627\u0633\u0631 \u0628\u0631 \u0622\u0632\u0645\u0648\u0646 "
        "\u062a\u062c\u0631\u0628\u06cc \u062a\u0623\u06a9\u06cc\u062f \u06a9\u0631\u062f."
    )

    canonical, report = _canonicalize_chunk_inline_originals(memory, chunk, initial)

    assert "(Jupp Esser)" in canonical
    assert report["changed"] is True
    assert report["before_hash"] != report["after_hash"]


def test_unresolved_critique_does_not_teach_style_but_remains_continuity_data() -> None:
    scores = {
        "accuracy": 9,
        "fluency": 9,
        "terminology": 9,
        "register": 9,
        "average": 9,
    }
    db = _EventDatabase(
        [
            {"event_type": "chunk_started", "payload": {}},
            {
                "event_type": "critique_completed",
                "payload": {
                    "valid": True,
                    "blocking_issue_count": 0,
                    "issue_count": 2,
                    "scores": scores,
                },
            },
        ]
    )

    assert not _chunk_style_approved(db, "job", 0)


def test_structure_audit_accepts_persian_enumeration_and_exposes_source_anomaly() -> None:
    source = "There are three differences. First, scope; second, method; third, time."
    candidate = (
        "\u0633\u0647 \u062a\u0641\u0627\u0648\u062a \u0648\u062c\u0648\u062f \u062f\u0627\u0631\u062f: "
        "\u0646\u062e\u0633\u062a\u060c \u062f\u0627\u0645\u0646\u0647\u061b \u062f\u0648\u0645\u060c "
        "\u0631\u0648\u0634\u061b \u0633\u0648\u0645\u060c \u0632\u0645\u0627\u0646."
    )
    assert audit_structure(source, candidate) == []
    anomaly = audit_structure(
        "It addresses two issues. First one; second two; third three.",
        "\u0627\u06cc\u0646 \u0645\u062a\u0646 \u0628\u0647 \u062f\u0648 \u0645\u0633\u0626\u0644\u0647 \u0645\u06cc\u200c\u067e\u0631\u062f\u0627\u0632\u062f. "
        "\u0646\u062e\u0633\u062a \u06cc\u06a9\u061b \u062f\u0648\u0645 \u062f\u0648\u061b \u0633\u0648\u0645 \u0633\u0647.",
    )
    assert [finding.classification for finding in anomaly] == [SOURCE_ANOMALY_PRESERVED]
    assert anomaly[0].details["source_excerpt"].startswith("It addresses")
    assert (
        anomaly[0].details["candidate_excerpt"].startswith("\u0627\u06cc\u0646 \u0645\u062a\u0646")
    )
    assert enumeration_length("Chapter 2; Chapter 3; Chapter 4") == 3


def test_catalog_lines_are_not_foreign_prose_leaks() -> None:
    assert unexpected_latin_prose("Jessop, Bob.", "Jessop, Bob.") == []
    assert (
        unexpected_latin_prose(
            "The state: past, present, future / Bob Jessop.",
            "The state: past, present, future / Bob Jessop.",
        )
        == []
    )
    assert unexpected_latin_prose(
        "This argument concerns the state.",
        "This argument concerns the state.",
    )


def test_source_confirmed_superscript_marker_reaches_docx() -> None:
    markers = _line_superscript_markers(
        [{"text": "3", "flags": 1}],
        "follows,3 historical",
    )
    assert markers and markers[0]["text"] == "3"
    metadata = {"superscript_markers": [{"text": "3", "relative_position": 0.5}]}
    text = "\u0627\u062f\u0639\u0627\u06cc \u0633\u0648\u0645 3 \u0627\u0633\u062a."
    assert source_superscript_spans(text, metadata) == [(10, 11)]

    translated = TranslatedDocument(
        paragraphs=[
            TranslatedParagraph(
                index=0,
                source_text="The third claim.3",
                translated_text=text,
                metadata=metadata,
            )
        ]
    )
    with TemporaryDirectory() as directory:
        path = Path(directory) / "marker.docx"
        DocxExporter().export(translated, path)
        document = Document(path)
        superscript = [
            run.text
            for paragraph in document.paragraphs
            for run in paragraph.runs
            if run.font.superscript
        ]
    assert "3" in superscript

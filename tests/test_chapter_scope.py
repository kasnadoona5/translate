from __future__ import annotations

import tempfile
import unittest
from pathlib import Path
from unittest.mock import MagicMock, patch

from docx import Document as DocxDocument

from tarjomeh.chunking.chunker import SemanticChunker
from tarjomeh.core.config import TarjomehConfig
from tarjomeh.core.pipeline import (
    TranslationPipeline,
    apply_chapter_selection,
    build_chapter_manifest,
)
from tarjomeh.jobs.database import JobDatabase, JobStatus
from tarjomeh.parsers.base import Chapter, Document, Paragraph, Section


def _document() -> Document:
    return Document(
        title="Test Book",
        chapters=[
            Chapter(
                title="Chapter 1: Opening",
                number=1,
                sections=[Section("", 2, [Paragraph("First chapter text.")])],
                metadata={"start_page": 1, "end_page": 4},
            ),
            Chapter(
                title="Chapter 2: Argument",
                number=2,
                sections=[Section("", 2, [Paragraph("Second chapter text.")])],
                metadata={"start_page": 5, "end_page": 9},
            ),
        ],
    )


class TestChapterStructure(unittest.TestCase):
    def test_manifest_and_chunks_preserve_original_chapter_identity(self) -> None:
        document = _document()
        manifest = build_chapter_manifest(document)
        selected = apply_chapter_selection(document, [2])
        chunks = SemanticChunker(max_tokens=100).chunk(selected)

        self.assertEqual([item["position"] for item in manifest], [1, 2])
        self.assertEqual(len(document.chapters), 2)
        self.assertEqual(manifest[1]["start_page"], 5)
        self.assertEqual(len(chunks), 1)
        self.assertEqual(chunks[0].metadata["chapter_position"], 2)
        self.assertEqual(chunks[0].metadata["chapter_number"], 2)

    def test_chunk_metadata_is_persisted_and_loaded(self) -> None:
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as temp_dir:
            db = JobDatabase(Path(temp_dir) / "jobs.db")
            db.create_job("chapter-meta", "book.txt", {})
            chunks = SemanticChunker(max_tokens=100).chunk(_document())
            db.save_chunks("chapter-meta", chunks)

            loaded = db.get_chunks("chapter-meta")

            self.assertEqual(loaded[0]["metadata"]["chapter_position"], 1)
            self.assertEqual(loaded[1]["metadata"]["chapter_position"], 2)


class TestChapterPipeline(unittest.TestCase):
    def _config(self) -> TarjomehConfig:
        config = TarjomehConfig()
        config.output.format = "txt"
        config.output.bilingual_mode = "target_only"
        config.translation.enable_critique = False
        config.translation.enable_back_translation = False
        config.translation.enable_web_context = False
        config.translation.enable_book_research = False
        config.glossary.enable_auto_extraction = False
        config.memory.enable_4layer = False
        return config

    @staticmethod
    def _fake_translation(**kwargs: object) -> str:
        chunk = kwargs["chunk"]
        return "FA " + chunk.text

    def test_selected_chapter_is_a_complete_scoped_job(self) -> None:
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as temp_dir:
            root = Path(temp_dir)
            source = root / "book.txt"
            source.write_text(
                "Chapter 1: Opening\nFirst chapter text.\n\n"
                "Chapter 2: Argument\nSecond chapter text.",
                encoding="utf-8",
            )
            output = root / "selected.txt"
            config = self._config()
            config.translation.chapter_selection = [2]
            pipeline = TranslationPipeline(config)
            pipeline.db = JobDatabase(root / "jobs.db")
            pipeline._translate_single_chunk = MagicMock(
                side_effect=self._fake_translation
            )

            pipeline.run(source, output, job_id="selected-job")

            job = pipeline.db.get_job("selected-job")
            self.assertEqual(job["raw_status"], JobStatus.COMPLETED)
            self.assertEqual(len(pipeline.db.get_chunks("selected-job")), 1)
            rendered = output.read_text(encoding="utf-8")
            self.assertIn("Second chapter text", rendered)
            self.assertNotIn("First chapter text", rendered)

    def test_checkpoint_exports_prefix_then_resume_finishes_book(self) -> None:
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as temp_dir:
            root = Path(temp_dir)
            source = root / "book.txt"
            source.write_text(
                "Chapter 1: Opening\nFirst chapter text.\n\n"
                "Chapter 2: Argument\nSecond chapter text.",
                encoding="utf-8",
            )
            output = root / "translated.txt"
            config = self._config()
            config.translation.stop_after_chapter = 1
            pipeline = TranslationPipeline(config)
            pipeline.db = JobDatabase(root / "jobs.db")
            pipeline._translate_single_chunk = MagicMock(
                side_effect=self._fake_translation
            )

            pipeline.run(source, output, job_id="checkpoint-job")

            paused = pipeline.db.get_job("checkpoint-job")
            self.assertEqual(paused["raw_status"], JobStatus.PAUSED)
            preview = output.read_text(encoding="utf-8")
            self.assertIn("First chapter text", preview)
            self.assertNotIn("Second chapter text", preview)

            pipeline.run(source, output, job_id="checkpoint-job")

            completed = pipeline.db.get_job("checkpoint-job")
            self.assertEqual(completed["raw_status"], JobStatus.COMPLETED)
            final = output.read_text(encoding="utf-8")
            self.assertIn("First chapter text", final)
            self.assertIn("Second chapter text", final)
            checkpoints = pipeline.db.get_job_artifact(
                "checkpoint-job", "chapter_checkpoints"
            )
            self.assertEqual(checkpoints["reached_positions"], [1])

    def test_checkpoint_exports_partial_docx_then_resume_finishes_book(self) -> None:
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as temp_dir:
            root = Path(temp_dir)
            source = root / "book.txt"
            source.write_text(
                "Chapter 1: Opening\nFirst chapter text.\n\n"
                "Chapter 2: Argument\nSecond chapter text.",
                encoding="utf-8",
            )
            output = root / "translated.docx"
            config = self._config()
            config.output.format = "docx"
            config.translation.stop_after_chapter = 1
            pipeline = TranslationPipeline(config)
            pipeline.db = JobDatabase(root / "jobs.db")
            pipeline._translate_single_chunk = MagicMock(
                side_effect=self._fake_translation
            )

            result = pipeline.run(source, output, job_id="checkpoint-docx-job")

            paused = pipeline.db.get_job("checkpoint-docx-job")
            self.assertEqual(paused["raw_status"], JobStatus.PAUSED)
            self.assertEqual(result.output_path, output)
            self.assertEqual(Path(paused["output_path"]), output)
            self.assertTrue(output.is_file())
            preview = "\n".join(
                paragraph.text for paragraph in DocxDocument(output).paragraphs
            )
            self.assertIn("First chapter text", preview)
            self.assertNotIn("Second chapter text", preview)

            pipeline.run(source, output, job_id="checkpoint-docx-job")

            completed = pipeline.db.get_job("checkpoint-docx-job")
            self.assertEqual(completed["raw_status"], JobStatus.COMPLETED)
            final = "\n".join(
                paragraph.text for paragraph in DocxDocument(output).paragraphs
            )
            self.assertIn("First chapter text", final)
            self.assertIn("Second chapter text", final)

    def test_checkpoint_is_not_visible_as_paused_before_preview_export(self) -> None:
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as temp_dir:
            root = Path(temp_dir)
            source = root / "book.txt"
            source.write_text(
                "Chapter 1: Opening\nFirst chapter text.\n\n"
                "Chapter 2: Argument\nSecond chapter text.",
                encoding="utf-8",
            )
            output = root / "translated.docx"
            config = self._config()
            config.output.format = "docx"
            config.translation.stop_after_chapter = 1
            pipeline = TranslationPipeline(config)
            pipeline.db = JobDatabase(root / "jobs.db")
            pipeline._translate_single_chunk = MagicMock(
                side_effect=self._fake_translation
            )
            observed_statuses: list[str] = []

            from tarjomeh.exporters.docx_exporter import DocxExporter

            class InspectingDocxExporter(DocxExporter):
                def export(self, *args: object, **kwargs: object) -> None:
                    job = pipeline.db.get_job("checkpoint-order-job")
                    observed_statuses.append(job["raw_status"])
                    super().export(*args, **kwargs)

            with patch(
                "tarjomeh.core.pipeline.get_exporter",
                return_value=InspectingDocxExporter,
            ):
                pipeline.run(source, output, job_id="checkpoint-order-job")

            self.assertEqual(observed_statuses, [JobStatus.RUNNING])
            paused = pipeline.db.get_job("checkpoint-order-job")
            self.assertEqual(paused["raw_status"], JobStatus.PAUSED)
            self.assertEqual(Path(paused["output_path"]), output)
            self.assertTrue(output.is_file())

    def test_failed_checkpoint_export_is_retryable(self) -> None:
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as temp_dir:
            root = Path(temp_dir)
            source = root / "book.txt"
            source.write_text(
                "Chapter 1: Opening\nFirst chapter text.\n\n"
                "Chapter 2: Argument\nSecond chapter text.",
                encoding="utf-8",
            )
            output = root / "translated.docx"
            config = self._config()
            config.output.format = "docx"
            config.translation.stop_after_chapter = 1
            pipeline = TranslationPipeline(config)
            pipeline.db = JobDatabase(root / "jobs.db")
            pipeline._translate_single_chunk = MagicMock(
                side_effect=self._fake_translation
            )

            with patch(
                "tarjomeh.core.pipeline.get_exporter",
                side_effect=RuntimeError("preview export failed"),
            ):
                with self.assertRaisesRegex(RuntimeError, "preview export failed"):
                    pipeline.run(source, output, job_id="checkpoint-retry-job")

            checkpoints = pipeline.db.get_job_artifact(
                "checkpoint-retry-job", "chapter_checkpoints"
            ) or {}
            self.assertEqual(checkpoints.get("reached_positions", []), [])
            failed = pipeline.db.get_job("checkpoint-retry-job")
            self.assertEqual(failed["raw_status"], JobStatus.PAUSED_ERROR)
            self.assertIn(
                "Chapter checkpoint preview export failed",
                failed["error_message"],
            )

    def test_checkpoint_resume_restores_web_search_state(self) -> None:
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as temp_dir:
            root = Path(temp_dir)
            source = root / "book.txt"
            source.write_text(
                "Chapter 1: Opening\nFirst chapter text.\n\n"
                "Chapter 2: Argument\nSecond chapter text.",
                encoding="utf-8",
            )
            output = root / "translated.txt"
            config = self._config()
            config.translation.stop_after_chapter = 1
            pipeline = TranslationPipeline(config)
            pipeline.db = JobDatabase(root / "jobs.db")
            pipeline._translate_single_chunk = MagicMock(
                side_effect=self._fake_translation
            )
            first_searcher = MagicMock()
            first_searcher.export_state.return_value = {
                "provider": {"queries_used": 7, "cache": {}},
            }
            resumed_searcher = MagicMock()
            resumed_searcher.export_state.return_value = {
                "provider": {"queries_used": 7, "cache": {}},
            }

            with patch(
                "tarjomeh.core.pipeline.WebContextSearcher",
                side_effect=[first_searcher, resumed_searcher],
            ):
                pipeline.run(source, output, job_id="search-resume-job")
                saved_state = pipeline.db.get_job_artifact(
                    "search-resume-job", "web_search_state"
                )
                pipeline.run(source, output, job_id="search-resume-job")

            resumed_searcher.import_state.assert_called_once_with(saved_state)


if __name__ == "__main__":
    unittest.main()

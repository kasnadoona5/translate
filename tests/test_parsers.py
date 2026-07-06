"""Unit tests for Phase 3: Document Parsers."""

from __future__ import annotations

import tempfile
import unittest
from pathlib import Path
from unittest.mock import MagicMock, patch

import sys

# Try importing real packages. If they fail, mock them to prevent ModuleNotFoundError on import.
try:
    import ebooklib
    from ebooklib import epub
except ImportError:
    mock_ebooklib = MagicMock()
    mock_ebooklib.ITEM_DOCUMENT = 9
    mock_epub = MagicMock()
    mock_ebooklib.epub = mock_epub
    sys.modules["ebooklib"] = mock_ebooklib
    sys.modules["ebooklib.epub"] = mock_epub
    ebooklib = mock_ebooklib
    epub = mock_epub

try:
    import fitz
except ImportError:
    mock_fitz = MagicMock()
    mock_fitz.TEXT_PRESERVE_WHITESPACE = 1
    class MockRect:
        def __init__(self, x0, y0, x1, y1):
            self.x0 = x0
            self.y0 = y0
            self.x1 = x1
            self.y1 = y1
            self.height = y1 - y0
    mock_fitz.Rect = MockRect
    sys.modules["fitz"] = mock_fitz

try:
    import docx
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph as DocxParagraph
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False
    mock_docx = MagicMock()
    mock_ns = MagicMock()
    mock_ns.qn = lambda x: x
    mock_table = MagicMock()
    mock_para = MagicMock()
    sys.modules["docx"] = mock_docx
    sys.modules["docx.oxml.ns"] = mock_ns
    sys.modules["docx.table"] = mock_table
    sys.modules["docx.text.paragraph"] = mock_para

from tarjomeh.parsers.base import Document, Chapter, Section, Paragraph
from tarjomeh.parsers.txt_parser import TxtParser
from tarjomeh.parsers.markdown_parser import MarkdownParser
from tarjomeh.parsers.srt_parser import SrtParser
from tarjomeh.parsers.docx_parser import DocxParser
from tarjomeh.parsers.epub_parser import EpubParser
from tarjomeh.parsers.pdf_parser import PyMuPDFParser, DocLayoutParser
from tarjomeh.parsers.ocr_preprocessor import OCRPreprocessor


class TestTxtParser(unittest.TestCase):
    """Test plain-text file parser and heuristics."""

    def test_txt_parser_heuristics(self) -> None:
        content = (
            "Chapter 1\n"
            "This is paragraph one.\n"
            "Line two of paragraph one.\n"
            "\n"
            "This is paragraph two.\n"
            "\n"
            "\n"
            "CHAPTER II: Next Part\n"
            "This is chapter two text.\n"
        )
        with tempfile.NamedTemporaryFile("w", suffix=".txt", delete=False, encoding="utf-8") as f:
            f.write(content)
            temp_path = Path(f.name)

        try:
            parser = TxtParser()
            doc = parser.parse(temp_path)

            self.assertEqual(doc.metadata["format_type"], "txt")
            self.assertEqual(doc.title, "Chapter 1")
            self.assertEqual(len(doc.chapters), 2)

            ch1 = doc.chapters[0]
            self.assertEqual(ch1.title, "Chapter 1")
            self.assertEqual(ch1.number, 1)
            self.assertEqual(len(ch1.sections), 1)
            self.assertEqual(len(ch1.sections[0].paragraphs), 2)
            self.assertEqual(ch1.sections[0].paragraphs[0].text, "This is paragraph one. Line two of paragraph one.")
            self.assertEqual(ch1.sections[0].paragraphs[1].text, "This is paragraph two.")

            ch2 = doc.chapters[1]
            self.assertEqual(ch2.title, "Next Part")
            self.assertEqual(ch2.number, 2)
            self.assertEqual(len(ch2.sections[0].paragraphs), 1)
            self.assertEqual(ch2.sections[0].paragraphs[0].text, "This is chapter two text.")
        finally:
            temp_path.unlink()

    def test_txt_parser_fallback(self) -> None:
        content = "Line one.\n\nLine two.\n"
        with tempfile.NamedTemporaryFile("w", suffix=".txt", delete=False, encoding="utf-8") as f:
            f.write(content)
            temp_path = Path(f.name)

        try:
            parser = TxtParser()
            doc = parser.parse(temp_path)
            self.assertEqual(len(doc.chapters), 1)
            self.assertEqual(doc.chapters[0].title, "Document")
            self.assertEqual(len(doc.chapters[0].sections[0].paragraphs), 2)
        finally:
            temp_path.unlink()


class TestMarkdownParser(unittest.TestCase):
    """Test Markdown parser heading levels and frontmatter."""

    def test_markdown_parser(self) -> None:
        content = (
            "---\n"
            "title: The Republic\n"
            "author: Plato\n"
            "---\n"
            "\n"
            "# Book I\n"
            "Some introductory text.\n"
            "\n"
            "## Section A\n"
            "Argument text.\n"
            "\n"
            "> This is a blockquote.\n"
            "\n"
            "```python\n"
            "print('code block')\n"
            "```\n"
            "\n"
            "### Level 3 Title\n"
            "Details here.\n"
        )
        with tempfile.NamedTemporaryFile("w", suffix=".md", delete=False, encoding="utf-8") as f:
            f.write(content)
            temp_path = Path(f.name)

        try:
            parser = MarkdownParser()
            doc = parser.parse(temp_path)

            self.assertEqual(doc.title, "The Republic")
            self.assertEqual(doc.metadata["author"], "Plato")
            self.assertEqual(len(doc.chapters), 1)

            ch = doc.chapters[0]
            self.assertEqual(ch.title, "Book I")
            
            # Should have 2 sections: 1. Default section (before ## Section A) and 2. Section A
            self.assertEqual(len(ch.sections), 2)
            
            # Default section
            self.assertEqual(len(ch.sections[0].paragraphs), 1)
            self.assertEqual(ch.sections[0].paragraphs[0].text, "Some introductory text.")

            # Section A
            sec = ch.sections[1]
            self.assertEqual(sec.title, "Section A")
            self.assertEqual(len(sec.paragraphs), 5)
            self.assertEqual(sec.paragraphs[0].text, "Argument text.")
            self.assertTrue(sec.paragraphs[1].metadata.get("blockquote"))
            self.assertTrue(sec.paragraphs[2].metadata.get("code_block"))
            self.assertEqual(sec.paragraphs[3].metadata.get("heading_level"), 3)
            self.assertEqual(sec.paragraphs[4].text, "Details here.")
        finally:
            temp_path.unlink()


class TestSrtParser(unittest.TestCase):
    """Test SRT parser timecodes and hour grouping."""

    def test_srt_parser(self) -> None:
        content = (
            "1\n"
            "00:00:01,123 --> 00:00:03,456\n"
            "Hello, how are you?\n"
            "\n"
            "2\n"
            "00:00:04,500 --> 00:00:07,000\n"
            "I am fine, thanks.\n"
            "\n"
            "3\n"
            "01:00:10,000 --> 01:00:12,123\n"
            "An hour later...\n"
        )
        with tempfile.NamedTemporaryFile("w", suffix=".srt", delete=False, encoding="utf-8") as f:
            f.write(content)
            temp_path = Path(f.name)

        try:
            parser = SrtParser()
            doc = parser.parse(temp_path)

            self.assertEqual(doc.metadata["format_type"], "srt")
            self.assertEqual(doc.metadata["cue_count"], 3)
            
            # Since cue 3 starts at 01:00:10 (Hour 1), it should split into 2 chapters
            self.assertEqual(len(doc.chapters), 2)
            self.assertEqual(doc.chapters[0].title, "Opening")
            self.assertEqual(doc.chapters[1].title, "Hour 1")

            p1 = doc.chapters[0].sections[0].paragraphs[0]
            self.assertEqual(p1.text, "Hello, how are you?")
            self.assertEqual(p1.metadata["index"], 1)
            self.assertEqual(p1.metadata["start_tc"], "00:00:01,123")
            self.assertEqual(p1.metadata["end_tc"], "00:00:03,456")
        finally:
            temp_path.unlink()


class TestDocxParser(unittest.TestCase):
    """Test Microsoft Word DOCX parser."""

    @unittest.skipUnless(HAS_DOCX, "python-docx is not installed")
    def test_docx_parser_with_docx(self) -> None:
        doc = docx.Document()
        doc.core_properties.title = "Sample Book"
        doc.core_properties.author = "Jane Austen"
        
        doc.add_heading("Chapter 1: Introduction", level=1)
        doc.add_paragraph("This is the first paragraph.")
        doc.add_heading("Section 1.1", level=2)
        doc.add_paragraph("This is section details.")
        
        # Add table
        t = doc.add_table(rows=2, cols=2)
        t.cell(0, 0).text = "A1"
        t.cell(0, 1).text = "B1"
        t.cell(1, 0).text = "A2"
        t.cell(1, 1).text = "B2"

        with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
            doc.save(f.name)
            temp_path = Path(f.name)

        try:
            parser = DocxParser()
            document = parser.parse(temp_path)

            self.assertEqual(document.title, "Sample Book")
            self.assertEqual(document.metadata["author"], "Jane Austen")
            self.assertEqual(len(document.chapters), 1)

            ch = document.chapters[0]
            self.assertEqual(ch.title, "Chapter 1: Introduction")
            self.assertEqual(len(ch.sections), 2)  # default + Section 1.1

            # Check table is serialized as pipe-delimited
            table_para = ch.sections[1].paragraphs[1]
            self.assertTrue(table_para.metadata.get("is_table"))
            self.assertIn("A1 | B1\nA2 | B2", table_para.text)
        finally:
            temp_path.unlink()

    @patch("tarjomeh.parsers.docx_parser.DocxDocument")
    def test_docx_parser_mocked(self, mock_docx_cls: MagicMock) -> None:
        mock_doc = MagicMock()
        mock_docx_cls.return_value = mock_doc
        
        mock_doc.core_properties.title = "Mocked Title"
        mock_doc.core_properties.author = "Mocked Author"
        
        # Create mock elements in XML body
        p1 = MagicMock()
        p1.tag = "p"
        
        p1_obj = MagicMock()
        p1_obj.text = "Hello Paragraph"
        p1_obj.style.name = "Normal"
        p1_obj._element = MagicMock()
        
        mock_doc.element.body = [p1]
        
        # Mocking DocxParagraph construction inside DocxParser
        with patch("tarjomeh.parsers.docx_parser.DocxParagraph", return_value=p1_obj):
            parser = DocxParser()
            with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tf:
                temp_path = Path(tf.name)
            try:
                doc = parser.parse(temp_path)
            finally:
                temp_path.unlink()
            
            self.assertEqual(doc.title, "Mocked Title")
            self.assertEqual(doc.metadata["author"], "Mocked Author")
            self.assertEqual(len(doc.chapters), 1)
            self.assertEqual(doc.chapters[0].sections[0].paragraphs[0].text, "Hello Paragraph")


class TestEpubParser(unittest.TestCase):
    """Test EPUB parser using mocks for ebooklib."""

    @patch("tarjomeh.parsers.epub_parser.epub.read_epub")
    def test_epub_parser(self, mock_read_epub: MagicMock) -> None:
        mock_book = MagicMock()
        mock_read_epub.return_value = mock_book

        # Mock metadata
        mock_book.get_metadata.side_effect = lambda namespace, name: {
            ("DC", "title"): ["My EPUB Title"],
            ("DC", "creator"): ["EPUB Author"],
            ("DC", "language"): ["en"],
        }.get((namespace, name), [])

        # Mock spine
        mock_book.spine = [("item_1", True), ("item_2", True)]

        # Mock items
        item1 = MagicMock()
        item1.get_id.return_value = "item_1"
        item1.get_type.return_value = ebooklib.ITEM_DOCUMENT
        item1.get_body_content.return_value = b"<h1>Chapter 1</h1><p>Paragraph in chapter 1.</p>"

        item2 = MagicMock()
        item2.get_id.return_value = "item_2"
        item2.get_type.return_value = ebooklib.ITEM_DOCUMENT
        item2.get_body_content.return_value = b"<h2>Section A</h2><p>Paragraph in Section A.</p>"

        mock_book.get_items.return_value = [item1, item2]
        mock_book.toc = []

        parser = EpubParser()
        with tempfile.NamedTemporaryFile(suffix=".epub", delete=False) as tf:
            temp_path = Path(tf.name)
        try:
            doc = parser.parse(temp_path)
        finally:
            temp_path.unlink()

        self.assertEqual(doc.title, "My EPUB Title")
        self.assertEqual(doc.metadata["author"], "EPUB Author")
        self.assertEqual(doc.metadata["language"], "en")
        self.assertEqual(len(doc.chapters), 2)

        self.assertEqual(doc.chapters[0].title, "Chapter 1")
        self.assertEqual(doc.chapters[0].sections[0].paragraphs[0].text, "Chapter 1")
        self.assertEqual(doc.chapters[0].sections[0].paragraphs[1].text, "Paragraph in chapter 1.")

        self.assertEqual(doc.chapters[1].title, "Section A")


class TestPdfParsers(unittest.TestCase):
    """Test PyMuPDFParser and DocLayoutParser."""

    @patch("tarjomeh.parsers.pdf_parser.fitz.open")
    def test_pymupdf_parser_with_toc(self, mock_fitz_open: MagicMock) -> None:
        mock_doc = MagicMock()
        mock_fitz_open.return_value = mock_doc

        mock_doc.name = "sample.pdf"
        mock_doc.metadata = {"title": "PDF Title", "author": "PDF Author"}
        mock_doc.page_count = 5
        mock_doc.get_toc.return_value = [
            [1, "Chapter 1", 1],
            [1, "Chapter 2", 3],
        ]

        # Mock page blocks
        mock_page = MagicMock()
        mock_doc.__getitem__.return_value = mock_page
        mock_page.rect.height = 800
        mock_page.get_images.return_value = []
        mock_page.get_text.side_effect = lambda opt, **kwargs: (
            "Some body text" if opt == "text" else {
                "blocks": [
                    {
                        "type": 0,
                        "lines": [
                            {
                                "spans": [
                                    {"text": "Text in block", "size": 10.0, "font": "Arial"}
                                ]
                            }
                        ],
                        "bbox": [50, 100, 200, 120]
                    }
                ]
            }
        )

        parser = PyMuPDFParser()
        with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as tf:
            temp_path = Path(tf.name)
        try:
            doc = parser.parse(temp_path)
        finally:
            temp_path.unlink()

        self.assertEqual(doc.title, "PDF Title")
        self.assertEqual(doc.metadata["author"], "PDF Author")
        self.assertEqual(len(doc.chapters), 2)
        self.assertEqual(doc.chapters[0].title, "Chapter 1")
        self.assertEqual(doc.chapters[1].title, "Chapter 2")

    @patch("tarjomeh.parsers.pdf_parser.DocLayoutParser._load_model")
    @patch("tarjomeh.parsers.pdf_parser.PyMuPDFParser.parse")
    def test_doclayout_parser_fallback(self, mock_pymupdf_parse: MagicMock, mock_load_model: MagicMock) -> None:
        mock_load_model.return_value = False
        parser = DocLayoutParser()
        with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as tf:
            temp_path = Path(tf.name)
        try:
            parser.parse(temp_path)
        finally:
            temp_path.unlink()
        
        # Should fallback to PyMuPDFParser
        mock_pymupdf_parse.assert_called_once()


class TestOCRPreprocessor(unittest.TestCase):
    """Test OCR preprocessor initialization and exception handling."""

    @patch("tarjomeh.parsers.ocr_preprocessor.OCRPreprocessor._ensure_ocrmypdf")
    def test_ocr_preprocessor_success(self, mock_ensure: MagicMock) -> None:
        mock_ocrmypdf = MagicMock()
        mock_ensure.return_value = mock_ocrmypdf
        
        def mock_ocr(input_file, output_file, **kwargs):
            Path(output_file).write_bytes(b"dummy pdf content")
            return 0
        mock_ocrmypdf.ocr.side_effect = mock_ocr

        # We need a dummy file that exists
        with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as f:
            temp_path = Path(f.name)

        try:
            preprocessor = OCRPreprocessor(languages="eng+fas")
            out_path = preprocessor.preprocess(temp_path)
            
            # Verify ocr method was called
            mock_ocrmypdf.ocr.assert_called_once()
            self.assertEqual(out_path.name, f"{temp_path.stem}_ocr.pdf")
        finally:
            # Clean up output file if it exists, plus the input file
            temp_path.unlink()
            if 'out_path' in locals() and out_path.exists():
                out_path.unlink()
                # If it created a temp directory, it will be cleaned up on process exit or we can delete it
                if out_path.parent.name.startswith("tarjomeh_ocr_"):
                    out_path.parent.rmdir()


if __name__ == "__main__":
    unittest.main()

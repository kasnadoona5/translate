"""Explicit saved-prefix exports retain canonical text and all render gates."""
import copy
from pathlib import Path
from unittest.mock import MagicMock

import docx
import pytest

from tarjomeh.chunking.chunker import Chunk
from tarjomeh.core.config import TarjomehConfig
from tarjomeh.core.pipeline import TranslationPipeline, _canonical_chunk_paragraph_identity
from tarjomeh.jobs.database import JobDatabase
from tarjomeh.jobs.output_coverage import read_output_coverage
from tarjomeh.parsers.base import Chapter, Document, Paragraph, Section

FA = ['\u0627\u06cc\u0646 \u0645\u062a\u0646 \u0646\u062e\u0633\u062a \u0627\u0633\u062a.', '\u0627\u06cc\u0646 \u0645\u062a\u0646 \u062f\u0648\u0645 \u0627\u0633\u062a.']


@pytest.fixture
def saved(tmp_path):
    db = JobDatabase(tmp_path / 'jobs.db')
    source = tmp_path / 'source.txt'
    source.write_text('First source.\n\nSecond source.\n\nLater source.', encoding='utf-8')
    config = TarjomehConfig()
    config.output.format = 'docx'
    config.output.term_notes = 'none'
    chunks = [Chunk(index=i, text=t, chapter_title='Section', section_title='', metadata={'paragraph_indices': [i]})
              for i, t in enumerate(('First source.', 'Second source.', 'Later source.'))]
    document = Document(title='Synthetic book', chapters=[Chapter(title='Section', sections=[Section(
        title='', level=1, paragraphs=[Paragraph(c.text) for c in chunks])])])
    db.create_job('sample', source, config.to_dict())
    db.save_chunks('sample', chunks)
    identities = {}
    for i, target in enumerate(FA):
        db.update_chunk('sample', i, 'needs_review' if i == 1 else 'completed', target)
        _, identities[str(i)] = _canonical_chunk_paragraph_identity(chunks[i], target)
    db.save_job_artifact('sample', 'canonical_chunk_paragraphs_v1', {'version': 1, 'chunks': identities})
    db.save_memory_state('sample', {'proper_nouns': {}, 'bilingual_summary': {'english_summary': 'Unchanged'}})
    db.update_job_status('sample', 'paused')
    pipeline = object.__new__(TranslationPipeline)
    pipeline.config, pipeline.db = config, db
    pipeline.worker_id = 'export-worker'
    pipeline.current_job_id, pipeline._worker_claimed = None, False
    pipeline._parse_and_chunk = MagicMock(return_value=(document, chunks))
    pipeline._start_lease_heartbeat = MagicMock()
    pipeline._stop_lease_heartbeat = MagicMock()
    pipeline.llm_client = MagicMock()
    return pipeline, chunks


@pytest.mark.parametrize('mode', ['target_only', 'inline', 'side_by_side'])
def test_saved_prefix_exports_without_translation_or_memory_change(saved, mode):
    pipeline, chunks = saved
    pipeline.config.output.bilingual_mode = mode
    db = pipeline.db
    before = copy.deepcopy((db.get_chunks('sample'), db.get_memory_state('sample')))
    result = pipeline.export_completed_job('sample', saved_prefix=True)
    assert result.name.endswith('_PARTIAL_REVIEW.docx')
    assert result.is_file()
    text = '\n'.join(p.text for p in docx.Document(result).paragraphs)
    assert 'Later source' not in text
    assert 'PARTIAL / REVIEW' in docx.Document(result).sections[0].header.paragraphs[0].text
    assert before == (db.get_chunks('sample'), db.get_memory_state('sample'))
    assert db.get_job('sample')['status'] == 'paused'
    assert db.get_worker_lease('sample')['state'] == 'released'
    assert not pipeline.llm_client.mock_calls
    coverage = read_output_coverage(db, db.get_job('sample'))
    assert coverage['published_finished_count'] == 2
    assert coverage['unpublished_finished_count'] == 0
    assert coverage['publication_kind'] == 'saved_prefix_review'
    assert coverage['total_selected_chunks'] == 3
    assert db.get_job_artifact('sample', 'render_identity_audit')['passed']


def test_default_full_export_still_requires_complete_range(saved):
    pipeline, _ = saved
    with pytest.raises(RuntimeError, match='missing completed translation'):
        pipeline.export_completed_job('sample')


@pytest.mark.parametrize('defect', ['gap', 'empty', 'source', 'indices', 'identity', 'busy', 'running', 'no_prefix'])
def test_saved_prefix_refuses_unproved_or_busy_scope(saved, defect):
    pipeline, chunks = saved
    db = pipeline.db
    if defect == 'gap':
        db.update_chunk('sample', 0, 'pending')
    elif defect == 'empty':
        db.update_chunk('sample', 0, 'completed', '')
    elif defect == 'source':
        chunks[0].text = 'Changed source.'
    elif defect == 'indices':
        chunks[0].metadata['paragraph_indices'] = [1]
    elif defect == 'identity':
        db.save_job_artifact('sample', 'canonical_chunk_paragraphs_v1', {})
    elif defect == 'busy':
        db.claim_worker('sample', 'other-worker')
    elif defect == 'running':
        db.update_job_status('sample', 'running')
    else:
        for i in range(2):
            db.update_chunk('sample', i, 'pending')
    before = copy.deepcopy((db.get_chunks('sample'), db.get_memory_state('sample')))
    with pytest.raises((ValueError, RuntimeError)):
        pipeline.export_completed_job('sample', saved_prefix=True)
    assert before == (db.get_chunks('sample'), db.get_memory_state('sample'))
    assert not list(Path(db.get_job('sample')['input_path']).parent.glob('*PARTIAL_REVIEW.docx'))


def test_saved_prefix_tamper_cannot_publish_or_overwrite_previous_file(saved, monkeypatch):
    pipeline, _ = saved
    path = Path(pipeline.db.get_job('sample')['input_path']).with_name('protected.docx')
    path.write_bytes(b'previous publication')
    original = pipeline._export_with_render_identity
    def tamper(job, document, exporter, output, ledger, fmt):
        document.paragraphs[0].translated_text += ' UNEXPLAINED'
        return original(job, document, exporter, output, ledger, fmt)
    monkeypatch.setattr(pipeline, '_export_with_render_identity', tamper)
    with pytest.raises(RuntimeError, match='Export blocked'):
        pipeline.export_completed_job('sample', path, saved_prefix=True)
    assert path.read_bytes() == b'previous publication'
    assert pipeline.db.get_worker_lease('sample')['state'] == 'released'


def test_saved_prefix_does_not_silently_change_format_or_selection(saved):
    pipeline, _ = saved
    with pytest.raises(ValueError):
        pipeline.export_completed_job('sample', saved_prefix=True, output_format='epub')
    with pytest.raises(ValueError):
        pipeline.export_completed_job('sample', saved_prefix=True, chapter_positions=[1])


def test_saved_prefix_preserves_chapter_checkpoint_progress(saved):
    pipeline, _ = saved
    db = pipeline.db
    state = {'version': 3, 'reached_positions': [1], 'pending': {'chapter_position': 2}, 'latest_position': 1}
    db.save_job_artifact('sample', 'chapter_checkpoints', state)
    pipeline.export_completed_job('sample', saved_prefix=True)
    after = db.get_job_artifact('sample', 'chapter_checkpoints')
    assert all(after[k] == v for k,v in state.items())
    assert after['latest_publication']['boundary_chunk_index'] == 1


def test_export_does_not_reclaim_stale_active_worker(saved):
    import sqlite3
    pipeline, _ = saved
    db = pipeline.db
    db.claim_worker('sample', 'old-owner')
    with db._get_connection() as conn:
        conn.execute("UPDATE job_workers SET heartbeat_at='2000-01-01' WHERE job_id='sample'")
    result = db.claim_worker('sample', 'exporter', allow_stale_reclaim=False)
    assert result['acquired'] is False
    assert db.get_worker_lease('sample')['worker_id'] == 'old-owner'


def test_protected_publication_identity_survives_render_failure(saved, monkeypatch):
    pipeline, _ = saved
    db = pipeline.db
    old_identity = {'old': 'identity'}
    old_checkpoint = {'version': 3, 'latest_publication': {'old': 'publication'}}
    db.save_job_artifact('sample', 'canonical_document_identity', old_identity)
    db.save_job_artifact('sample', 'chapter_checkpoints', old_checkpoint)
    def fail(*args):
        raise RuntimeError('Export blocked: injected read-back mismatch')
    monkeypatch.setattr(pipeline, '_export_with_render_identity', fail)
    with pytest.raises(RuntimeError, match='Export blocked'):
        pipeline.export_completed_job('sample', saved_prefix=True)
    assert db.get_job_artifact('sample', 'canonical_document_identity') == old_identity
    assert db.get_job_artifact('sample', 'chapter_checkpoints') == old_checkpoint
    assert db.get_job('sample')['output_path'] is None


def test_saved_prefix_rejects_non_docx_output_path(saved):
    pipeline, _ = saved
    with pytest.raises(ValueError, match='separate regular DOCX'):
        pipeline.export_completed_job('sample', Path('wrong.txt'), saved_prefix=True)


def test_saved_prefix_never_overwrites_source_docx(saved):
    pipeline, _ = saved
    old = Path(pipeline.db.get_job('sample')['input_path'])
    source = old.with_suffix('.docx')
    source.write_bytes(old.read_bytes())
    with pipeline.db._get_connection() as conn:
        conn.execute('UPDATE jobs SET input_path=? WHERE id=?', (str(source), 'sample'))
    before = source.read_bytes()
    with pytest.raises(ValueError, match='separate regular DOCX'):
        pipeline.export_completed_job('sample', source, saved_prefix=True)
    assert source.read_bytes() == before


def test_saved_prefix_rejects_changed_source_file(saved):
    pipeline, _ = saved
    Path(pipeline.db.get_job('sample')['input_path']).write_text('Changed upload.', encoding='utf-8')
    with pytest.raises(ValueError, match='source file changed'):
        pipeline.export_completed_job('sample', saved_prefix=True)


def test_saved_prefix_requires_one_verified_target_unit_per_source_paragraph(saved):
    pipeline, _ = saved
    db = pipeline.db
    identities = db.get_job_artifact('sample', 'canonical_chunk_paragraphs_v1')
    identities['chunks']['0']['target_offsets'] = []
    identities['chunks']['0']['target_hashes'] = []
    db.save_job_artifact('sample', 'canonical_chunk_paragraphs_v1', identities)
    with pytest.raises(ValueError, match='exact saved paragraph identities'):
        pipeline.export_completed_job('sample', saved_prefix=True)

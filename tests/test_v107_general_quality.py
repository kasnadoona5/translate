"""General structure, memory, and language regressions for v10.7."""

from __future__ import annotations

from pathlib import Path

from docx import Document

from tarjomeh.core.pipeline import _anchor_only_repair_is_valid
from tarjomeh.exporters.base import TranslatedDocument, TranslatedParagraph
from tarjomeh.exporters.docx_exporter import DocxExporter, contents_display_title
from tarjomeh.memory.proper_nouns import (
    has_exact_observed_anchor,
    observed_bilingual_target,
)
from tarjomeh.parsers.pdf_parser import (
    PyMuPDFParser,
    _contents_entry_metadata,
    _page_looks_like_contents,
    _split_contents_entry_lines,
)
from tarjomeh.quality.integrity import repair_source_grounded_language_artifacts


def _line(text: str, y: float) -> dict:
    return {"text": text, "bbox": (50.0, y, 360.0, y + 10.0)}


def test_contents_geometry_preserves_rows_and_wrapped_title() -> None:
    blocks = [
        {"lines": [_line("Contents", 100.0)]},
        {"lines": [
            _line("Preface", 200.0),
            _line("viii", 200.0),
            _line("1 Introduction", 220.0),
            _line("1", 220.0),
            _line("9 Liberal Democracy, Exceptional States,", 240.0),
            _line("and the New Normal", 250.0),
            _line("211", 250.0),
            _line("References", 270.0),
            _line("257", 270.0),
        ]},
    ]

    assert _page_looks_like_contents(blocks)
    groups = _split_contents_entry_lines(blocks[1]["lines"])
    assert len(groups) == 4
    wrapped = _contents_entry_metadata(groups[2])
    assert wrapped == {
        "structure_role": "contents_entry",
        "toc_page_label": "211",
        "toc_source_title": (
            "9 Liberal Democracy, Exceptional States, and the New Normal"
        ),
        "toc_entry_kind": "numbered",
        "toc_level": 1,
    }


def test_pdf_parser_defaults_to_structure_version_four() -> None:
    assert PyMuPDFParser().structure_version == 4


def test_docx_contents_export_uses_borderless_aligned_rows(tmp_path: Path) -> None:
    document = TranslatedDocument(
        title="Test",
        paragraphs=[
            TranslatedParagraph(
                index=0,
                source_text="Contents",
                translated_text="\u0641\u0647\u0631\u0633\u062a \u0645\u0637\u0627\u0644\u0628",
                heading_level=1,
                metadata={"structure_role": "heading"},
            ),
            TranslatedParagraph(
                index=1,
                source_text="Preface viii",
                translated_text="\u067e\u06cc\u0634\u06af\u0641\u062a\u0627\u0631 viii",
                metadata={
                    "structure_role": "contents_entry",
                    "toc_page_label": "viii",
                    "toc_entry_kind": "unnumbered",
                    "toc_level": 0,
                },
            ),
            TranslatedParagraph(
                index=2,
                source_text="1 Introduction 1",
                translated_text="\u06f1 \u0645\u0642\u062f\u0645\u0647 \u06f1",
                metadata={
                    "structure_role": "contents_entry",
                    "toc_page_label": "1",
                    "toc_entry_kind": "numbered",
                    "toc_level": 1,
                },
            ),
        ],
    )
    output = tmp_path / "contents.docx"

    DocxExporter().export(document, output, "target_only")

    exported = Document(output)
    assert len(exported.tables) == 1
    assert len(exported.tables[0].rows) == 2
    assert exported.tables[0].cell(0, 0).text == "\u067e\u06cc\u0634\u06af\u0641\u062a\u0627\u0631"
    assert exported.tables[0].cell(0, 1).text == "viii"
    assert exported.tables[0].cell(1, 0).text == "\u06f1 \u0645\u0642\u062f\u0645\u0647"
    assert exported.tables[0].cell(1, 1).text == "1"


def test_contents_display_title_does_not_strip_non_page_content() -> None:
    metadata = {"toc_page_label": "53"}
    assert contents_display_title("\u0641\u0635\u0644 \u0633\u0648\u0645 53", metadata) == "\u0641\u0635\u0644 \u0633\u0648\u0645"
    assert contents_display_title("\u0641\u0635\u0644 \u0633\u0648\u0645", metadata) == "\u0641\u0635\u0644 \u0633\u0648\u0645"


def test_observed_name_alignment_handles_diacritics_and_hyphens() -> None:
    assert observed_bilingual_target(
        "\u0648 \u0698\u0627\u0646-\u0698\u0627\u06a9 \u0631\u0648\u0633\u0648 (Jean-Jacques Rousseau)",
        "Jean-Jacques Rousseau",
        "person",
    ) == "\u0698\u0627\u0646-\u0698\u0627\u06a9 \u0631\u0648\u0633\u0648"
    assert observed_bilingual_target(
        "\u0627\u0650\u0645\u0631\u06cc\u0634 \u062f\u0648 \u0648\u0627\u062a\u0644 (Emmerich de Vattel)",
        "Emmerich de Vattel",
        "person",
    ) == "\u0627\u0650\u0645\u0631\u06cc\u0634 \u062f\u0648 \u0648\u0627\u062a\u0644"


def test_uncertain_coordinated_entity_is_not_admitted() -> None:
    assert observed_bilingual_target(
        "\u0627\u0642\u06cc\u0627\u0646\u0648\u0633 \u0627\u0637\u0644\u0633 \u0634\u0645\u0627\u0644\u06cc \u0648 \u0645\u0646\u0637\u0642\u0647 \u06cc\u0648\u0631\u0648 "
        "(North Atlantic and Eurozone)",
        "North Atlantic and Eurozone",
        "source_grounded_entity",
    ) == ""


def test_legal_anchor_accepts_source_grounded_year_suffix() -> None:
    assert has_exact_observed_anchor(
        "\u0642\u0627\u0646\u0648\u0646 \u0646\u0645\u0648\u0646\u0647 (Sample Copyright Act 1988)",
        "Sample Copyright Act",
    )


def test_anchor_repair_rejects_any_non_insertion_edit() -> None:
    before = "\u06cc\u0648\u067e \u0627\u0633\u0631 \u0628\u0631 \u0622\u0632\u0645\u0648\u0646 \u062a\u062c\u0631\u0628\u06cc \u062a\u0627\u06a9\u06cc\u062f \u06a9\u0631\u062f."
    after = "\u06cc\u0648\u067e \u0627\u0633\u0631 (Jupp Esser) \u0628\u0631 \u0622\u0632\u0648\u0646 \u062a\u062c\u0631\u0628\u06cc \u062a\u0627\u06a9\u06cc\u062f \u06a9\u0631\u062f."
    assert not _anchor_only_repair_is_valid(
        before, after, ["Jupp Esser"], {"Jupp Esser": "person"}
    )
    valid = "\u06cc\u0648\u067e \u0627\u0633\u0631 (Jupp Esser) \u0628\u0631 \u0622\u0632\u0645\u0648\u0646 \u062a\u062c\u0631\u0628\u06cc \u062a\u0627\u06a9\u06cc\u062f \u06a9\u0631\u062f."
    assert _anchor_only_repair_is_valid(
        before, valid, ["Jupp Esser"], {"Jupp Esser": "person"}
    )


def test_source_grounded_language_repair_is_narrow() -> None:
    source = "The account studies the very longue duree and history of the state."
    source = source.replace("longue duree", "longue dur\u00e9e")
    target = (
        "\u0627\u06cc\u0646 \u062a\u0628\u06cc\u06cc\u0646 longue dur\u00e9e \u0648 \u062a\u0627\u0631\u06cc\u062e \u062a\u0627\u0631\u06cc\u062e "
        "\u062f\u0648\u0644\u062a \u0631\u0627 \u0628\u0631\u0631\u0633\u06cc \u0645\u06cc\u200c\u06a9\u0646\u062f."
    )

    repaired, report = repair_source_grounded_language_artifacts(source, target)

    assert "(longue dur\u00e9e)" in repaired
    assert "\u062a\u0627\u0631\u06cc\u062e \u062a\u0627\u0631\u06cc\u062e" not in repaired
    assert report["repair_count"] == 2


def test_source_repetition_is_preserved() -> None:
    target = "\u0628\u0633\u06cc\u0627\u0631 \u0628\u0633\u06cc\u0627\u0631 \u0645\u0647\u0645 \u0627\u0633\u062a."
    repaired, report = repair_source_grounded_language_artifacts(
        "It is very very important.", target
    )
    assert repaired == target
    assert report["repair_count"] == 0

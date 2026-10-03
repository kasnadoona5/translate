"""Read-only evidence must not grant authority or hide failures."""

import hashlib
import io
from unittest.mock import MagicMock

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from tarjomeh.core.evidence_audit import extra_final_refinement_evidence
from tarjomeh.core.render_identity import docx_contents_evidence
from tarjomeh.exporters.base import TranslatedDocument, TranslatedParagraph
from tarjomeh.exporters.docx_exporter import DocxExporter
from tarjomeh.glossary.book_review import book_term_review_evidence


@pytest.mark.parametrize("value", ["2", "-1", "1.5", "true"])
def test_upload_rejects_invalid_extra_allowance_before_reserving_worker(
    tmp_path, monkeypatch, value,
):
    from tarjomeh.web import app as module

    monkeypatch.chdir(tmp_path)
    monkeypatch.setenv("UI_SECRET_TOKEN", "local-test")
    submit = MagicMock()
    monkeypatch.setattr(module._executor, "submit", submit)
    client = module.create_app().test_client()
    response = client.post("/api/translate", headers={"Authorization": "Bearer local-test"}, data={
        "file": (io.BytesIO(b"A source."), "book.txt"), "extra_final_refine_attempts": value,
    })
    assert response.status_code == 400
    submit.assert_not_called()
    assert not list((tmp_path / "jobs/uploads").iterdir())


def test_attempt_accounting_retains_obsolete_failures_and_physical_retries():
    context = {"stage": "extra_final_refinement"}
    events = [
        {"chunk_index": 1, "event_type": "chunk_started", "payload": {}},
        {"chunk_index": 1, "event_type": "llm_call_attempt", "payload": {
            "quality_attempt_context": context, "success": False,
            "operation": "refinement", "duration_seconds": 2,
        }},
        {"chunk_index": 1, "event_type": "chunk_started", "payload": {}},
        {"chunk_index": 1, "event_type": "llm_call_attempt", "payload": {
            "quality_attempt_context": context, "success": True,
            "operation": "refinement", "response_model": "test-model",
            "completion_tokens": 12, "duration_seconds": 3,
        }},
        {"chunk_index": 1, "event_type": "llm_call_attempt", "payload": {
            "success": True, "operation": "translation",
        }},
    ]
    report = extra_final_refinement_evidence(
        {"entries": {"1": {"consumed": 1, "state": "finished"}}}, events, 1
    )
    assert report["active"]["physical_attempts"] == 1
    assert report["lifetime"]["physical_attempts"] == 2
    assert report["lifetime"]["failures"] == 1
    assert report["lifetime"]["llm_seconds"] == 5
    assert report["active"]["served_models"] == {"test-model": 1}
    assert report["consumed_chunk_count"] == 1


@pytest.mark.parametrize("value", [0, 1])
def test_upload_forwards_explicit_allowance_without_running_llm(tmp_path, monkeypatch, value):
    from tarjomeh.web import app as module

    monkeypatch.chdir(tmp_path)
    monkeypatch.setenv("UI_SECRET_TOKEN", "local-test")
    monkeypatch.setattr(module, "_active_jobs", {})
    monkeypatch.setattr(module, "_job_worker_ids", {})
    monkeypatch.setattr(module, "_progress_queues", {})
    submit = MagicMock()
    monkeypatch.setattr(module._executor, "submit", submit)
    client = module.create_app().test_client()
    response = client.post("/api/translate", headers={"Authorization": "Bearer local-test"}, data={
        "file": (io.BytesIO(b"A source."), "book.txt"), "extra_final_refine_attempts": str(value),
    })
    assert response.status_code == 202
    run_job = submit.call_args.args[0]
    captured = dict(zip(run_job.__code__.co_freevars,
                        [cell.cell_contents for cell in run_job.__closure__], strict=True))
    assert captured["config_overrides"]["translation.extra_final_refine_attempts"] == value


def test_corrupt_or_unfinished_accounting_is_never_pass():
    assert extra_final_refinement_evidence({"entries": []}, [], 1)["status"] == "FAIL"
    assert extra_final_refinement_evidence(
        {"entries": {"0": {"consumed": 2}}}, [], 1
    )["status"] == "FAIL"
    assert extra_final_refinement_evidence(
        {"entries": {"0": {"consumed": 1, "state": "reserved"}}}, [], 1
    )["status"] == "REVIEW"


def test_proposal_approval_is_distinct_from_ticks_scope_and_lexical_presence():
    source = "A concept is discussed."
    sha = hashlib.sha256(source.encode()).hexdigest()
    proposals = [{"status": "approved", "source": "concept", "target": "concept"}]
    report = book_term_review_evidence({"proposals": proposals}, [], [])
    assert report["approved_proposal_count"] == 1
    assert report["ticked_occurrence_count"] == report["scoped_paragraph_count"] == 0
    scope = {"matched": [{"paragraph_index": 0, "source_paragraph_hash": sha,
                           "source": "concept", "target": "concept"}]}
    events = [{"chunk_index": 0, "event_type": "book_term_scope_resolved", "payload": scope}]
    chunks = [{"chunk_index": 0, "status": "completed", "text": source,
               "translation": source, "metadata": {"paragraph_indices": [0]}}]
    report = book_term_review_evidence({"proposals": proposals}, events, chunks)
    assert report["lexically_present_paragraph_term_count"] == 1
    assert report["paragraph_term_presence"][0]["semantic_accuracy"] == (
        "requires_source_based_human_review"
    )
    chunks[0]["metadata"] = "invalid JSON"
    report = book_term_review_evidence({"proposals": proposals}, events, chunks)
    assert report["paragraph_term_presence"][0]["presence"] == "REVIEW"


def test_native_contents_units_are_read_from_the_actual_docx(tmp_path):
    document = TranslatedDocument(title="Book", author="", metadata={}, paragraphs=[
        TranslatedParagraph(index=0, source_text="Contents", translated_text="Contents",
                            metadata={"structure_role": "heading"}),
        *[TranslatedParagraph(index=i, source_text=f"Section {i} {i + 20}",
                              translated_text=f"Section {i} {i + 20}", metadata={
                                  "structure_role": "contents_entry", "toc_page_label": str(i + 20),
                              }) for i in range(1, 21)],
        TranslatedParagraph(index=21, source_text="Body.", translated_text="Body.",
                            metadata={"structure_role": "body"}),
    ])
    path = tmp_path / "contents.docx"
    DocxExporter({}).export(document=document, output_path=path, bilingual_mode="target_only")
    report = docx_contents_evidence(path, 20)
    assert report["status"] == "PASS"
    assert report["matched_native_rtl_rows"] is True
    assert report["native_page_tables"][0]["rows"] == 20
    assert report["native_page_tables"][0]["fixed_layout"]
    actual = Document(path)
    property_ = actual.tables[0]._tbl.tblPr.find(qn("w:bidiVisual"))
    if property_ is None:
        property_ = OxmlElement("w:bidiVisual")
        actual.tables[0]._tbl.tblPr.append(property_)
    property_.set(qn("w:val"), "0")
    actual.save(path)
    assert docx_contents_evidence(path, 20)["status"] == "REVIEW"

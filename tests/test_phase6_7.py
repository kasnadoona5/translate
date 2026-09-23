from __future__ import annotations

import gc
import json
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import AsyncMock, MagicMock, patch

from tarjomeh.context.book_researcher import BookResearcher
from tarjomeh.context.search_providers import (
    BaseSearchProvider,
    BraveSearchProvider,
    SearchProviderChain,
    SearchResult,
    TavilySearchProvider,
)
from tarjomeh.core.config import TarjomehConfig
from tarjomeh.core.llm_client import TruncatedCompletionError
from tarjomeh.core.pipeline import _research_context_for_memory
from tarjomeh.core.term_notes import (
    apply_term_notes,
    effective_term_notes_mode,
    ensure_inline_proper_noun_originals,
)
from tarjomeh.exporters.base import TranslatedDocument, TranslatedParagraph
from tarjomeh.exporters.markdown_exporter import MarkdownExporter
from tarjomeh.exporters.docx_exporter import DocxExporter, HAS_DOCX
from tarjomeh.exporters.epub_exporter import EpubExporter
from tarjomeh.glossary.manager import GlossaryManager
from tarjomeh.jobs.database import JobDatabase
from tarjomeh.memory.manager import MemoryManager
from tarjomeh.persian.typography import PersianTypographer


class TestPhase6TermNotes(unittest.TestCase):
    def setUp(self) -> None:
        self.glossary = GlossaryManager()
        self.glossary.add_term("capital", "سرمایه")
        self.typographer = PersianTypographer({
            "convert_numerals": False,
            "normalize_zwnj": False,
            "fix_punctuation": False,
        })
        self.document = TranslatedDocument(
            title="Test",
            paragraphs=[
                TranslatedParagraph(
                    index=0,
                    source_text="Capital and Gramsci appear here.",
                    translated_text="سرمایه و گرامشی در اینجا ظاهر می‌شوند.",
                ),
                TranslatedParagraph(
                    index=1,
                    source_text="Capital appears again.",
                    translated_text="سرمایه دوباره ظاهر می‌شود.",
                ),
            ],
        )

    def test_defaults_preserve_existing_inline_behaviour(self) -> None:
        config = TarjomehConfig()
        self.assertEqual(config.output.term_notes, "inline")
        self.assertFalse(config.translation.enable_book_research)
        self.assertEqual(
            apply_term_notes(
                self.document,
                self.glossary,
                {"Gramsci": "گرامشی"},
                self.typographer,
                mode="inline",
            ),
            [],
        )
        self.assertNotIn("term_notes", self.document.metadata)

    def test_unsupported_note_format_falls_back_to_inline(self) -> None:
        self.assertEqual(effective_term_notes_mode("footnote", "pdf"), "inline")
        self.assertEqual(effective_term_notes_mode("footnote", "docx"), "footnote")

    def test_refinement_cannot_remove_inline_proper_noun_original(self) -> None:
        document = TranslatedDocument(
            title="Test",
            paragraphs=[TranslatedParagraph(
                index=0,
                source_text="Intacta rr2 Pro is discussed.",
                translated_text="«اینتکتا آرآر ۲ پرو» مطرح می‌شود.",
            )],
        )
        restored = ensure_inline_proper_noun_originals(
            document,
            {"Intacta rr2 Pro": "اینتکتا آرآر ۲ پرو"},
            self.typographer,
        )
        self.assertEqual(restored, 1)
        self.assertIn("» (Intacta rr2 Pro)", document.paragraphs[0].translated_text)
        self.assertEqual(
            ensure_inline_proper_noun_originals(
                document,
                {"Intacta rr2 Pro": "اینتکتا آرآر ۲ پرو"},
                self.typographer,
            ),
            0,
        )
    def test_note_mode_removes_inline_parenthetical_memory_instruction(self) -> None:
        config = TarjomehConfig()
        config.output.term_notes = "footnote"
        config.output.format = "docx"
        memory = MemoryManager(config)
        memory.proper_nouns.add_noun("Gramsci", "گرامشی")
        context = memory.get_context_for_chunk(
            type("ChunkStub", (), {"text": "Gramsci"})()
        ).proper_nouns
        self.assertIn("exporter adds", context)
        self.assertNotIn("add (Gramsci) after", context)

    def test_notes_are_numbered_once_in_document_order(self) -> None:
        notes = apply_term_notes(
            self.document,
            self.glossary,
            {"Gramsci": "گرامشی"},
            self.typographer,
            mode="footnote",
        )
        self.assertEqual(
            [(note["number"], note["original"]) for note in notes],
            [(1, "capital"), (2, "Gramsci")],
        )
        self.assertEqual(
            [ref["number"] for ref in self.document.paragraphs[0].metadata["term_note_refs"]],
            [1, 2],
        )
        self.assertNotIn("term_note_refs", self.document.paragraphs[1].metadata)

    def test_markdown_export_has_native_notes(self) -> None:
        apply_term_notes(
            self.document,
            self.glossary,
            {"Gramsci": "گرامشی"},
            self.typographer,
            mode="footnote",
        )
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "book.md"
            MarkdownExporter().export(self.document, path)
            content = path.read_text(encoding="utf-8")
        self.assertIn("سرمایه[^1]", content)
        self.assertIn("گرامشی[^2]", content)
        self.assertIn("[^1]: capital (سرمایه)", content)
        self.assertEqual(content.count("[^1]:"), 1)

    def test_epub_export_has_linked_footnotes(self) -> None:
        apply_term_notes(
            self.document,
            self.glossary,
            {"Gramsci": "گرامشی"},
            self.typographer,
            mode="footnote",
        )
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "book.epub"
            EpubExporter().export(self.document, path)
            with zipfile.ZipFile(path) as archive:
                chapter = archive.read("OEBPS/chapter_1.xhtml").decode()
                notes = archive.read("OEBPS/notes.xhtml").decode()
        self.assertIn('epub:type="noteref"', chapter)
        self.assertIn('id="term-note-1"', notes)
        self.assertIn("capital", notes)

    @unittest.skipUnless(HAS_DOCX, "python-docx is not installed")
    def test_docx_export_uses_reliable_endnotes_section(self) -> None:
        import docx

        apply_term_notes(
            self.document,
            self.glossary,
            {"Gramsci": "گرامشی"},
            self.typographer,
            mode="footnote",
        )
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "book.docx"
            DocxExporter().export(self.document, path)
            text = chr(10).join(p.text for p in docx.Document(path).paragraphs)
        self.assertIn("یادداشت‌ها", text)
        self.assertIn("capital", text)


class TestPhase7Persistence(unittest.TestCase):
    def test_unapproved_research_is_explicitly_non_mandatory_context(self) -> None:
        context = _research_context_for_memory({
            "status": "completed",
            "book_context": "Book context",
            "terms": [
                {
                    "source": "capital",
                    "target": "پایتخت",
                    "status": "suggested",
                    "identity_supported": True,
                    "term_supported": True,
                },
                {
                    "source": "state",
                    "target": "کشور",
                    "status": "rejected",
                },
            ],
        })
        self.assertIn("not mandatory terminology", context)
        self.assertIn("capital -> پایتخت", context)
        self.assertNotIn("state -> کشور", context)

    def test_job_artifact_round_trip(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            db = JobDatabase(Path(tmp) / "jobs.db")
            db.create_job("job-1", "book.pdf", {})
            payload = {
                "status": "completed",
                "terms": [{"source": "capital", "target": "سرمایه"}],
            }
            db.save_job_artifact("job-1", "book_research", payload)
            self.assertEqual(
                db.get_job_artifact("job-1", "book_research"),
                payload,
            )
            del db
            gc.collect()

    def test_book_context_memory_is_additive_and_serialized(self) -> None:
        memory = MemoryManager(TarjomehConfig())
        memory.book_context = "Political theory context"
        state = memory.to_dict()
        self.assertIn("proper_nouns", state)
        self.assertIn("bilingual_summary", state)
        self.assertIn("past_translations", state)
        self.assertIn("short_term_context", state)

        restored = MemoryManager(TarjomehConfig())
        restored.from_dict(state)
        self.assertEqual(restored.book_context, "Political theory context")
        self.assertIn(
            "Political theory context",
            restored.get_context_for_chunk(
                type("ChunkStub", (), {"text": ""})()
            ).format(),
        )


class TestBookResearcher(unittest.IsolatedAsyncioTestCase):
    def test_follow_up_query_objects_use_query_text_not_python_repr(self) -> None:
        queries = BookResearcher._normalise_follow_ups(
            [
                {"query": "author concept", "rationale": "resolve ambiguity"},
                {"query": "author concept", "rationale": "duplicate"},
                {"rationale": "missing query"},
            ],
            existing=[],
            limit=3,
        )
        self.assertEqual(queries, ["author concept"])

    async def test_research_is_bounded_reviewable_and_source_checked(self) -> None:
        config = TarjomehConfig()
        llm = MagicMock()
        llm.chat = AsyncMock(return_value=json.dumps({
            "book_context": "A study of political operations.",
            "terms": [{
                "source": "operation",
                "target": "عملیات",
                "reason": "Central book concept",
                "confidence": "high",
                "source_urls": [
                    "https://example.test/book",
                    "https://invented.test/not-allowed",
                ],
            }],
        }))
        researcher = BookResearcher(config, llm)
        researcher.provider = MagicMock()
        researcher.provider.search = AsyncMock(return_value=[
            SearchResult(
                title="The Politics of Operations - Book page",
                url="https://example.test/book",
                snippet="The Politics of Operations is a book about operations.",
            )
        ])
        document = type("DocumentStub", (), {
            "title": "The Politics of Operations",
            "author": "Test Author",
            "raw_toc": ["Introduction"],
            "all_paragraphs": [
                type("ParagraphStub", (), {"text": "Operations shape politics."})()
            ],
        })()

        result = await researcher.research(document)

        self.assertEqual(result.status, "completed")
        self.assertEqual(len(result.terms), 1)
        self.assertTrue(result.terms[0]["is_auto"])
        self.assertEqual(result.terms[0]["status"], "suggested")
        self.assertEqual(
            result.terms[0]["source_urls"],
            ["https://example.test/book"],
        )
        self.assertEqual(researcher.provider.search.await_count, 5)

    async def test_research_uses_bounded_adaptive_follow_up_queries(self) -> None:
        config = TarjomehConfig()
        config.web_search.phase7_max_queries = 8
        llm = MagicMock()
        llm.chat = AsyncMock(side_effect=[
            json.dumps({
                "book_context": "Initial",
                "terms": [],
                "follow_up_queries": [
                    '"operative surface" author-specific meaning',
                    '"hitting the ground" political economy',
                    '"operative surface" author-specific meaning',
                    "variegation capital geography",
                    "ignored fourth query",
                ],
            }),
            json.dumps({
                "book_context": "Final evidence-grounded context",
                "terms": [],
                "follow_up_queries": [],
            }),
        ])
        researcher = BookResearcher(config, llm)
        researcher.provider = MagicMock()
        researcher.provider.search = AsyncMock(return_value=[
            SearchResult("Result", "https://example.test/result", "Evidence")
        ])
        document = type("DocumentStub", (), {
            "title": "Book",
            "author": "Author",
            "raw_toc": [],
            "all_paragraphs": [
                type("ParagraphStub", (), {"text": "Book excerpt."})()
            ],
        })()

        result = await researcher.research(document)

        self.assertEqual(len(result.queries), 8)
        self.assertEqual(researcher.provider.search.await_count, 8)
        self.assertEqual(llm.chat.await_count, 2)
        self.assertEqual(
            result.book_context,
            "Final evidence-grounded context",
        )

    async def test_research_without_external_sources_is_degraded(self) -> None:
        config = TarjomehConfig()
        llm = MagicMock()
        llm.chat = AsyncMock(return_value=json.dumps({
            "book_context": "Excerpt-only context",
            "terms": [],
            "follow_up_queries": [],
        }))
        researcher = BookResearcher(config, llm)
        researcher.provider = MagicMock()
        researcher.provider.search = AsyncMock(return_value=[])
        document = type("DocumentStub", (), {
            "title": "Book",
            "author": "",
            "raw_toc": [],
            "all_paragraphs": [
                type("ParagraphStub", (), {"text": "Book excerpt."})()
            ],
        })()

        result = await researcher.research(document)

        self.assertEqual(result.status, "degraded")
        self.assertEqual(result.book_context, "Excerpt-only context")

    async def test_research_recovers_with_bounded_evidence_batches(self) -> None:
        config = TarjomehConfig()
        llm = MagicMock()
        llm.chat = AsyncMock(side_effect=[
            TruncatedCompletionError("whole synthesis reached length"),
            json.dumps({
                "book_context": "Batch context",
                "terms": [{
                    "source": "operation",
                    "target": "operation-fa",
                    "source_urls": ["https://example.test/book"],
                }],
            }),
            json.dumps({
                "book_context": "Consolidated context",
                "terms": [{
                    "source": "operation",
                    "target": "operation-fa",
                    "source_urls": ["https://example.test/book"],
                }],
                "follow_up_queries": [],
            }),
        ])
        researcher = BookResearcher(config, llm)
        researcher.provider = MagicMock()
        researcher.provider.diagnostics = [{
            "provider": "tavily", "status": "success",
        }]
        researcher.provider.search = AsyncMock(return_value=[
            SearchResult(
                "Book page",
                "https://example.test/book",
                "Evidence",
            )
        ])
        document = type("DocumentStub", (), {
            "title": "Book",
            "author": "Author",
            "raw_toc": [],
            "all_paragraphs": [
                type("ParagraphStub", (), {"text": "Book excerpt."})()
            ],
        })()

        result = await researcher.research(document)

        self.assertEqual(result.status, "degraded")
        self.assertEqual(result.book_context, "Consolidated context")
        self.assertEqual(len(result.terms), 1)
        self.assertIn("TruncatedCompletionError", result.error)
        self.assertEqual(
            [call.args[0] for call in llm.set_operation.call_args_list],
            [
                "book_research_initial",
                "book_research_batch",
                "book_research_synthesis",
            ],
        )


class TestSearchProviderChain(unittest.IsolatedAsyncioTestCase):
    async def test_tavily_parses_official_json_response(self) -> None:
        response = MagicMock()
        response.status_code = 200
        response.json.return_value = {
            "results": [{
                "title": "Academic source",
                "url": "https://example.test/source",
                "content": "Relevant research evidence",
            }]
        }
        client = AsyncMock()
        client.post.return_value = response
        context = MagicMock()
        context.__aenter__ = AsyncMock(return_value=client)
        context.__aexit__ = AsyncMock(return_value=False)

        with patch(
            "tarjomeh.context.search_providers.httpx.AsyncClient",
            return_value=context,
        ):
            results = await TavilySearchProvider("secret").search("query")

        self.assertEqual(results[0].title, "Academic source")
        self.assertEqual(results[0].snippet, "Relevant research evidence")
        request = client.post.await_args
        self.assertNotIn("secret", str(request.kwargs["json"]))

    async def test_brave_parses_official_json_response(self) -> None:
        response = MagicMock()
        response.status_code = 200
        response.json.return_value = {
            "web": {"results": [{
                "title": "Book review",
                "url": "https://example.test/review",
                "description": "A scholarly review",
            }]}
        }
        client = AsyncMock()
        client.get.return_value = response
        context = MagicMock()
        context.__aenter__ = AsyncMock(return_value=client)
        context.__aexit__ = AsyncMock(return_value=False)

        with patch(
            "tarjomeh.context.search_providers.httpx.AsyncClient",
            return_value=context,
        ):
            results = await BraveSearchProvider("secret").search("query")

        self.assertEqual(results[0].url, "https://example.test/review")
        self.assertEqual(results[0].snippet, "A scholarly review")

    async def test_fallback_cache_and_budget(self) -> None:
        empty = MagicMock(spec=BaseSearchProvider)
        empty.name = "primary"
        empty.available = True
        empty.last_status = 503
        empty.last_error = "temporary"
        empty.search = AsyncMock(return_value=[])

        fallback = MagicMock(spec=BaseSearchProvider)
        fallback.name = "fallback"
        fallback.available = True
        fallback.last_status = 200
        fallback.last_error = ""
        fallback.search = AsyncMock(return_value=[
            SearchResult("Title", "https://example.test", "Snippet")
        ])

        chain = SearchProviderChain(
            [empty, fallback],
            max_retries=0,
            query_budget=1,
        )
        first = await chain.search("academic term")
        cached = await chain.search("  ACADEMIC   TERM ")
        blocked = await chain.search("second term")

        self.assertEqual(len(first), 1)
        self.assertEqual(cached, first)
        self.assertEqual(blocked, [])
        self.assertEqual(chain.queries_used, 1)
        fallback.search.assert_awaited_once()
        self.assertEqual(chain.diagnostics[-1]["status"], "budget_exhausted")

    async def test_state_round_trip_preserves_paid_query_cache(self) -> None:
        provider = MagicMock(spec=BaseSearchProvider)
        provider.name = "provider"
        provider.available = True
        provider.last_status = 200
        provider.last_error = ""
        provider.search = AsyncMock(return_value=[
            SearchResult("Title", "https://example.test", "Snippet")
        ])
        original = SearchProviderChain([provider], query_budget=3)
        await original.search("query")

        restored = SearchProviderChain([provider], query_budget=3)
        restored.import_state(original.export_state())
        results = await restored.search("query")

        self.assertEqual(len(results), 1)
        self.assertEqual(restored.queries_used, 1)
        provider.search.assert_awaited_once()

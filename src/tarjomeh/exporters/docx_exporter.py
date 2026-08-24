"""DOCX document exporter for Tarjomeh."""

from __future__ import annotations

import logging
import re
import unicodedata
from pathlib import Path
from tarjomeh.exporters.base import BaseExporter, TranslatedDocument, BilingualMode
from tarjomeh.exporters.term_notes import document_term_notes, paragraph_note_parts

logger = logging.getLogger(__name__)

_LATIN_PARENTHETICAL_RE = re.compile(r"\([^()\n]*[A-Za-z][^()\n]*\)")
_PERSIAN_LETTER_RE = re.compile(
    r"[\u0621-\u063a\u0641-\u064a\u066e-\u06d3\u06fa-\u06ff]"
)
_DIGIT_EQUIVALENTS = {
    "0": "0\u0660\u06f0", "1": "1\u0661\u06f1", "2": "2\u0662\u06f2",
    "3": "3\u0663\u06f3", "4": "4\u0664\u06f4", "5": "5\u0665\u06f5",
    "6": "6\u0666\u06f6", "7": "7\u0667\u06f7", "8": "8\u0668\u06f8",
    "9": "9\u0669\u06f9",
}
_DIGIT_TO_ASCII = str.maketrans(
    "\u0660\u0661\u0662\u0663\u0664\u0665\u0666\u0667\u0668\u0669"
    "\u06f0\u06f1\u06f2\u06f3\u06f4\u06f5\u06f6\u06f7\u06f8\u06f9",
    "01234567890123456789",
)
_TOC_TRAILING_LABEL_RE = re.compile(
    r"(?:^|\s)(?:[ivxlcdm]{1,12}|[0-9\u0660-\u0669\u06f0-\u06f9]{1,4})"
    r"[\s\u060c\u061b,;]*$",
    re.IGNORECASE,
)


def contents_display_title(text: str, metadata: dict) -> str:
    """Remove only the duplicated source page label from a translated TOC row."""
    if not str(metadata.get("toc_page_label", "")).strip():
        return (text or "").strip()
    match = _TOC_TRAILING_LABEL_RE.search(text or "")
    if not match:
        return (text or "").strip()
    title = (text or "")[:match.start()].rstrip(" \t\u060c\u061b,;")
    return title or (text or "").strip()


def source_superscript_spans(
    text: str,
    metadata: dict,
) -> list[tuple[int, int]]:
    """Locate only source-confirmed superscript markers in translated text."""
    records = metadata.get("superscript_markers", []) or []
    if not isinstance(records, list):
        return []
    selected: list[tuple[int, int]] = []
    for record in records:
        if not isinstance(record, dict):
            continue
        raw = str(record.get("text", "")).strip()
        normalized = raw.translate(_DIGIT_TO_ASCII)
        if normalized.isdigit():
            pattern = "".join(
                f"[{_DIGIT_EQUIVALENTS[digit]}]" for digit in normalized
            )
            matcher = re.compile(rf"(?<!\d){pattern}(?!\d)")
        elif raw in {"¹", "²", "³", "⁰", "⁴", "⁵", "⁶", "⁷", "⁸", "⁹"}:
            matcher = re.compile(re.escape(raw))
        elif raw in {"*", "†", "‡"}:
            matcher = re.compile(re.escape(raw))
        else:
            continue
        candidates = [
            match.span() for match in matcher.finditer(text or "")
            if not any(
                match.start() < end and match.end() > start
                for start, end in selected
            )
        ]
        if not candidates:
            continue
        try:
            expected = float(record.get("relative_position", 0.0))
        except (TypeError, ValueError):
            expected = 0.0
        chosen = min(
            candidates,
            key=lambda span: abs(
                ((span[0] + span[1]) / 2) / max(1, len(text or "")) - expected
            ),
        )
        selected.append(chosen)
    return sorted(selected)


def _mixed_direction_parts(value: str) -> list[tuple[str, bool]]:
    """Split mixed citation content into stable RTL and LTR run spans."""
    parts: list[tuple[str, bool]] = []
    buffer = ""
    direction: bool | None = None
    for char in value:
        bidi = unicodedata.bidirectional(char)
        char_direction: bool | None
        if bidi in {"R", "AL", "AN"}:
            char_direction = True
        elif bidi in {"L", "EN"}:
            char_direction = False
        else:
            char_direction = None
        if char_direction is not None and direction is not None and char_direction != direction:
            if buffer:
                parts.append((buffer, direction))
            buffer = char
            direction = char_direction
        else:
            buffer += char
            if direction is None and char_direction is not None:
                direction = char_direction
    if buffer:
        parts.append((buffer, True if direction is None else direction))
    return parts


def directional_target_parts(text: str) -> list[tuple[str, bool]]:
    """Return text/run-direction pairs without changing visible content."""
    output: list[tuple[str, bool]] = []
    cursor = 0
    for match in _LATIN_PARENTHETICAL_RE.finditer(text or ""):
        if match.start() > cursor:
            output.append((text[cursor:match.start()], True))
        parenthetical = match.group()
        inner = parenthetical[1:-1]
        if _PERSIAN_LETTER_RE.search(inner):
            output.append(("(", True))
            output.extend(_mixed_direction_parts(inner))
            output.append((")", True))
        else:
            output.append((parenthetical, False))
        cursor = match.end()
    if cursor < len(text or ""):
        output.append((text[cursor:], True))
    return [(value, rtl) for value, rtl in output if value]

try:
    import docx
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.shared import Cm, Pt, RGBColor
    from docx.oxml.ns import qn
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False


class DocxExporter(BaseExporter):
    """DOCX exporter.

    Generates RTL-formatted DOCX files with optional bilingual layouts.
    """

    def export(
        self,
        document: TranslatedDocument,
        output_path: str | Path,
        bilingual_mode: BilingualMode = "target_only",
    ) -> Path:
        if not HAS_DOCX:
            raise ImportError(
                "python-docx is required for DOCX export. "
                "Install it with:  pip install python-docx"
            )

        out = Path(output_path)
        out.parent.mkdir(parents=True, exist_ok=True)

        doc = docx.Document()
        doc.core_properties.title = document.title
        doc.core_properties.author = document.author

        for section in doc.sections:
            section.top_margin = Cm(2.2)
            section.bottom_margin = Cm(2.2)
            section.left_margin = Cm(2.2)
            section.right_margin = Cm(2.2)

        normal_style = doc.styles["Normal"]
        normal_style.font.name = "Vazirmatn"
        normal_style.font.size = Pt(11)
        normal_style.paragraph_format.space_after = Pt(6)
        normal_style.paragraph_format.line_spacing = 1.15

        for level, size in ((1, 16), (2, 14), (3, 12)):
            heading_style = doc.styles[f"Heading {level}"]
            heading_style.font.name = "Vazirmatn"
            heading_style.font.size = Pt(size)
            heading_style.font.bold = True
            heading_style.font.color.rgb = RGBColor(0, 0, 0)
            heading_style.paragraph_format.keep_with_next = True

        chapter_page_breaks = bool(
            document.metadata.get("chapter_page_breaks", True)
        )
        seen_chapters: set[int] = set()

        def apply_chapter_break(p_obj, metadata: dict) -> None:
            if not chapter_page_breaks:
                return
            position = metadata.get("chapter_position")
            if not isinstance(position, int):
                return
            if position in seen_chapters:
                return
            if seen_chapters:
                p_obj.paragraph_format.page_break_before = True
            seen_chapters.add(position)

        def set_on_off(parent, tag: str, enabled: bool) -> None:
            element = parent.find(qn(tag))
            if element is None:
                element = OxmlElement(tag)
                parent.append(element)
            element.set(qn("w:val"), "1" if enabled else "0")

        def set_run_fonts(r_obj, *, rtl: bool) -> None:
            rPr = r_obj._r.get_or_add_rPr()
            set_on_off(rPr, "w:rtl", rtl)
            rFonts = rPr.find(qn("w:rFonts"))
            if rFonts is None:
                rFonts = OxmlElement("w:rFonts")
                rPr.append(rFonts)
            if rtl:
                rFonts.set(qn("w:cs"), "Vazirmatn")
                rFonts.set(qn("w:eastAsia"), "Vazirmatn")
                r_obj.font.name = "Vazirmatn"
            else:
                rFonts.set(qn("w:ascii"), "Times New Roman")
                rFonts.set(qn("w:hAnsi"), "Times New Roman")
                r_obj.font.name = "Times New Roman"
            r_obj.font.size = Pt(11)

        def make_paragraph_rtl(p_obj, *, heading: bool = False) -> None:
            p_obj.alignment = (
                WD_ALIGN_PARAGRAPH.RIGHT
                if heading else WD_ALIGN_PARAGRAPH.JUSTIFY
            )
            pPr = p_obj._element.get_or_add_pPr()
            set_on_off(pPr, "w:bidi", True)
            fmt = p_obj.paragraph_format
            fmt.line_spacing = 1.15
            fmt.space_after = Pt(6)
            if heading:
                fmt.keep_with_next = True
                fmt.space_before = Pt(10)
                fmt.space_after = Pt(4)
            else:
                fmt.first_line_indent = Cm(0.5)
                fmt.widow_control = True

        def add_target_runs(p_obj, text: str, metadata: dict) -> None:
            source_markers = source_superscript_spans(text, metadata)
            cursor = 0

            def add_text(value: str, absolute_start: int) -> None:
                local_cursor = 0
                local_markers = [
                    (start - absolute_start, end - absolute_start)
                    for start, end in source_markers
                    if absolute_start <= start < end <= absolute_start + len(value)
                ]
                for start, end in local_markers:
                    if start > local_cursor:
                        for part, rtl in directional_target_parts(
                            value[local_cursor:start]
                        ):
                            run = p_obj.add_run(part)
                            set_run_fonts(run, rtl=rtl)
                    marker = p_obj.add_run(value[start:end])
                    set_run_fonts(marker, rtl=True)
                    marker.font.superscript = True
                    local_cursor = end
                if local_cursor < len(value):
                    for part, rtl in directional_target_parts(value[local_cursor:]):
                        run = p_obj.add_run(part)
                        set_run_fonts(run, rtl=rtl)

            for segment, ref in paragraph_note_parts(text, metadata):
                if segment:
                    absolute_start = text.find(segment, cursor)
                    if absolute_start < 0:
                        absolute_start = cursor
                    add_text(segment, absolute_start)
                    cursor = absolute_start + len(segment)
                if ref is not None:
                    marker = p_obj.add_run(
                        str(ref.get("display_number", ref["number"]))
                    )
                    marker.font.superscript = True
                    set_run_fonts(marker, rtl=True)

        def apply_structural_format(p_obj, metadata: dict) -> None:
            if not metadata.get("is_table"):
                return
            # Preserve table-derived text as a compact reviewable block. Native
            # cells are emitted only when a parser provides an actual matrix.
            p_obj.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            p_obj.paragraph_format.first_line_indent = Cm(0)
            p_obj.paragraph_format.space_after = Pt(2)
            p_obj.paragraph_format.keep_together = True

        def add_contents_table(entries) -> None:
            """Emit preserved contents rows with stable RTL title/page alignment."""
            table = doc.add_table(rows=0, cols=2)
            section = doc.sections[-1]
            usable_width = int(section.page_width or 0) - int(
                section.left_margin or 0
            ) - int(section.right_margin or 0)
            total_twips = max(7200, int(usable_width / 635))
            page_twips = max(1100, int(total_twips * 0.14))
            title_twips = total_twips - page_twips

            table.alignment = WD_TABLE_ALIGNMENT.RIGHT
            table.autofit = False
            tbl_pr = table._tbl.tblPr
            for tag in ("w:bidiVisual", "w:tblW", "w:tblInd", "w:tblLayout"):
                for element in list(tbl_pr.findall(qn(tag))):
                    tbl_pr.remove(element)
            bidi_visual = OxmlElement("w:bidiVisual")
            bidi_visual.set(qn("w:val"), "1")
            tbl_pr.insert(0, bidi_visual)
            table_width = OxmlElement("w:tblW")
            table_width.set(qn("w:type"), "dxa")
            table_width.set(qn("w:w"), str(total_twips))
            tbl_pr.append(table_width)
            table_indent = OxmlElement("w:tblInd")
            table_indent.set(qn("w:type"), "dxa")
            table_indent.set(qn("w:w"), "0")
            tbl_pr.append(table_indent)
            fixed_layout = OxmlElement("w:tblLayout")
            fixed_layout.set(qn("w:type"), "fixed")
            tbl_pr.append(fixed_layout)

            grid = table._tbl.tblGrid
            for child in list(grid):
                grid.remove(child)
            for width in (title_twips, page_twips):
                grid_column = OxmlElement("w:gridCol")
                grid_column.set(qn("w:w"), str(width))
                grid.append(grid_column)
            borders = OxmlElement("w:tblBorders")
            for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
                border = OxmlElement(f"w:{edge}")
                border.set(qn("w:val"), "nil")
                borders.append(border)
            tbl_pr.append(borders)

            for entry in entries:
                row = table.add_row()
                row.height = Pt(15)
                row_pr = row._tr.get_or_add_trPr()
                row_pr.append(OxmlElement("w:cantSplit"))
                title_cell, page_cell = row.cells
                for cell, width in (
                    (title_cell, title_twips),
                    (page_cell, page_twips),
                ):
                    tc_pr = cell._tc.get_or_add_tcPr()
                    tc_width = tc_pr.find(qn("w:tcW"))
                    if tc_width is None:
                        tc_width = OxmlElement("w:tcW")
                        tc_pr.append(tc_width)
                    tc_width.set(qn("w:type"), "dxa")
                    tc_width.set(qn("w:w"), str(width))
                    cell_margin = OxmlElement("w:tcMar")
                    for edge in ("top", "start", "bottom", "end"):
                        margin = OxmlElement(f"w:{edge}")
                        margin.set(qn("w:type"), "dxa")
                        margin.set(qn("w:w"), "35" if edge in {"start", "end"} else "0")
                        cell_margin.append(margin)
                    tc_pr.append(cell_margin)
                page_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                title_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

                page_paragraph = page_cell.paragraphs[0]
                set_on_off(page_paragraph._p.get_or_add_pPr(), "w:bidi", False)
                page_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                page_paragraph.paragraph_format.space_after = Pt(0)
                page_run = page_paragraph.add_run(
                    str(entry.metadata.get("toc_page_label", ""))
                )
                set_run_fonts(page_run, rtl=False)

                title_paragraph = title_cell.paragraphs[0]
                make_paragraph_rtl(title_paragraph)
                title_paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                title_paragraph.paragraph_format.first_line_indent = Cm(0)
                title_paragraph.paragraph_format.space_after = Pt(0)
                if int(entry.metadata.get("toc_level", 0) or 0) > 0:
                    title_paragraph.paragraph_format.right_indent = Cm(0.45)
                display_title = contents_display_title(
                    entry.translated_text, entry.metadata
                )
                add_target_runs(title_paragraph, display_title, entry.metadata)
                if entry.metadata.get("toc_entry_kind") == "part":
                    for run in title_paragraph.runs:
                        run.bold = True

        if bilingual_mode == "side_by_side":
            table = doc.add_table(rows=0, cols=2)
            table.autofit = False
            
            for p in document.paragraphs:
                row = table.add_row()
                cell_en, cell_fa = row.cells
                
                # Left Column: English (LTR)
                p_en = cell_en.paragraphs[0]
                apply_chapter_break(p_en, p.metadata)
                if p.heading_level is not None:
                    p_en.style = doc.styles[f'Heading {min(p.heading_level, 9)}']
                run_en = p_en.add_run(p.source_text)
                
                # Right Column: Persian (RTL)
                p_fa = cell_fa.paragraphs[0]
                if p.heading_level is not None:
                    p_fa.style = doc.styles[f'Heading {min(p.heading_level, 9)}']
                add_target_runs(p_fa, p.translated_text, p.metadata)
                
                make_paragraph_rtl(p_fa, heading=p.heading_level is not None)
                apply_structural_format(p_fa, p.metadata)
                
        else:
            paragraph_index = 0
            while paragraph_index < len(document.paragraphs):
                p = document.paragraphs[paragraph_index]
                if (
                    bilingual_mode == "target_only"
                    and p.metadata.get("structure_role") == "contents_entry"
                ):
                    entries = []
                    while (
                        paragraph_index < len(document.paragraphs)
                        and document.paragraphs[paragraph_index].metadata.get(
                            "structure_role"
                        ) == "contents_entry"
                    ):
                        entries.append(document.paragraphs[paragraph_index])
                        paragraph_index += 1
                    add_contents_table(entries)
                    continue
                # Add Heading or Paragraph
                if bilingual_mode == "target_only":
                    if p.heading_level is not None:
                        p_fa = doc.add_heading(level=min(p.heading_level, 9))
                    else:
                        p_fa = doc.add_paragraph()
                    
                    add_target_runs(p_fa, p.translated_text, p.metadata)
                    make_paragraph_rtl(p_fa, heading=p.heading_level is not None)
                    apply_chapter_break(p_fa, p.metadata)
                    apply_structural_format(p_fa, p.metadata)
                    
                elif bilingual_mode == "inline":
                    # English paragraph (LTR)
                    if p.heading_level is not None:
                        p_en = doc.add_heading(level=min(p.heading_level, 9))
                    else:
                        p_en = doc.add_paragraph()
                    p_en.add_run(p.source_text)
                    apply_chapter_break(p_en, p.metadata)
                    
                    # Persian paragraph (RTL)
                    if p.heading_level is not None:
                        p_fa = doc.add_heading(level=min(p.heading_level, 9))
                    else:
                        p_fa = doc.add_paragraph()
                    add_target_runs(p_fa, p.translated_text, p.metadata)
                    make_paragraph_rtl(p_fa, heading=p.heading_level is not None)
                    apply_structural_format(p_fa, p.metadata)

                paragraph_index += 1

        notes = document_term_notes(document)
        if notes:
            heading = doc.add_heading("یادداشت‌ها", level=1)
            make_paragraph_rtl(heading, heading=True)
            for note in notes:
                note_para = doc.add_paragraph()
                note_run = note_para.add_run(
                    f"{note.get('display_number', note['number'])}. {note['original']} "
                    f"({note['transliteration']})"
                )
                make_paragraph_rtl(note_para)
                set_run_fonts(note_run, rtl=True)

        doc.save(str(out))
        return out

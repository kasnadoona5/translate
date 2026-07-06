"""DOCX document parser using *python-docx*.

Maps Word document structure onto the canonical
:class:`~tarjomeh.parsers.base.Document` model.  Heading 1 paragraphs
create chapter boundaries, Heading 2 creates section boundaries, and
footnotes / endnotes are preserved in paragraph metadata.
"""

from __future__ import annotations

import logging
from pathlib import Path
from typing import Any

from docx import Document as DocxDocument  # type: ignore[import-untyped]
from docx.oxml.ns import qn  # type: ignore[import-untyped]
from docx.table import Table  # type: ignore[import-untyped]
from docx.text.paragraph import Paragraph as DocxParagraph  # type: ignore[import-untyped]

from tarjomeh.parsers.base import (
    BaseParser,
    Chapter,
    Document,
    Paragraph,
    Section,
)

logger = logging.getLogger(__name__)

# ---------------------------------------------------------------------------
# Style helpers
# ---------------------------------------------------------------------------

_HEADING_STYLE_PREFIX = "Heading"


def _heading_level_from_style(style_name: str) -> int | None:
    """Extract heading level from a Word style name like ``'Heading 1'``.

    Returns *None* if *style_name* is not a heading style.
    """
    if not style_name:
        return None
    name = style_name.strip()
    if name.startswith(_HEADING_STYLE_PREFIX):
        suffix = name[len(_HEADING_STYLE_PREFIX) :].strip()
        if suffix.isdigit():
            return int(suffix)
    return None


# ---------------------------------------------------------------------------
# DocxParser
# ---------------------------------------------------------------------------


class DocxParser(BaseParser):
    """Parse Microsoft Word ``.docx`` files using *python-docx*.

    * **Heading 1** → chapter boundary
    * **Heading 2** → section boundary
    * **Heading 3–6** → paragraph with ``heading_level`` metadata
    * Footnotes and endnotes are attached as paragraph metadata.
    * Tables are serialized to a pipe-delimited text representation.
    """

    def parse(self, file_path: Path) -> Document:
        file_path = self._ensure_file(file_path)

        docx_doc = DocxDocument(str(file_path))

        title, author = self._extract_properties(docx_doc)
        footnotes = self._extract_footnotes(docx_doc)
        endnotes = self._extract_endnotes(docx_doc)

        chapters = self._build_chapters(docx_doc, footnotes, endnotes)

        return Document(
            title=title or file_path.stem,
            chapters=chapters,
            metadata={
                "author": author,
                "format_type": "docx",
                "language": "",
            },
        )

    # ------------------------------------------------------------------
    # Document properties
    # ------------------------------------------------------------------

    @staticmethod
    def _extract_properties(docx_doc: DocxDocument) -> tuple[str, str]:
        """Return ``(title, author)`` from core properties."""
        props = docx_doc.core_properties
        title = props.title or ""
        author = props.author or ""
        return title, author

    # ------------------------------------------------------------------
    # Footnotes & endnotes
    # ------------------------------------------------------------------

    @staticmethod
    def _extract_footnotes(docx_doc: DocxDocument) -> dict[str, str]:
        """Extract footnotes keyed by footnote ID.

        Uses raw XML traversal because *python-docx* does not expose
        footnotes through its public API.
        """
        notes: dict[str, str] = {}
        # Access the footnotes part if available
        try:
            footnotes_part = docx_doc.part.part_related_by(
                "http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes"
            )
        except KeyError:
            return notes
        except Exception as e:
            logger.warning("Error looking up footnotes part: %s", e)
            return notes

        if footnotes_part is None:
            return notes

        try:
            tree = footnotes_part.element
            for fn in tree.findall(qn("w:footnote")):
                fn_id = fn.get(qn("w:id"), "")
                if fn_id in ("-1", "0"):
                    # Separator and continuation-separator — skip
                    continue
                texts: list[str] = []
                for t_elem in fn.iter(qn("w:t")):
                    if t_elem.text:
                        texts.append(t_elem.text)
                if texts:
                    notes[fn_id] = " ".join(texts)
        except Exception:
            logger.debug("Failed to parse footnotes XML", exc_info=True)

        return notes

    @staticmethod
    def _extract_endnotes(docx_doc: DocxDocument) -> dict[str, str]:
        """Extract endnotes keyed by endnote ID (same approach as footnotes)."""
        notes: dict[str, str] = {}
        try:
            endnotes_part = docx_doc.part.part_related_by(
                "http://schemas.openxmlformats.org/officeDocument/2006/relationships/endnotes"
            )
        except KeyError:
            return notes
        except Exception as e:
            logger.warning("Error looking up endnotes part: %s", e)
            return notes

        if endnotes_part is None:
            return notes

        try:
            tree = endnotes_part.element
            for en in tree.findall(qn("w:endnote")):
                en_id = en.get(qn("w:id"), "")
                if en_id in ("-1", "0"):
                    continue
                texts: list[str] = []
                for t_elem in en.iter(qn("w:t")):
                    if t_elem.text:
                        texts.append(t_elem.text)
                if texts:
                    notes[en_id] = " ".join(texts)
        except Exception:
            logger.debug("Failed to parse endnotes XML", exc_info=True)

        return notes

    # ------------------------------------------------------------------
    # Footnote references within paragraphs
    # ------------------------------------------------------------------

    @staticmethod
    def _find_footnote_refs(para: DocxParagraph) -> list[str]:
        """Return footnote/endnote IDs referenced by *para*."""
        ids: list[str] = []
        for ref in para._element.iter(qn("w:footnoteReference")):
            fid = ref.get(qn("w:id"))
            if fid:
                ids.append(fid)
        for ref in para._element.iter(qn("w:endnoteReference")):
            eid = ref.get(qn("w:id"))
            if eid:
                ids.append(eid)
        return ids

    # ------------------------------------------------------------------
    # Table serialization
    # ------------------------------------------------------------------

    @staticmethod
    def _table_to_text(table: Table) -> str:
        """Convert a Word table to a simple pipe-delimited text string."""
        rows: list[str] = []
        for row in table.rows:
            cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            rows.append(" | ".join(cells))
        return "\n".join(rows)

    # ------------------------------------------------------------------
    # Build chapter structure
    # ------------------------------------------------------------------

    def _build_chapters(
        self,
        docx_doc: DocxDocument,
        footnotes: dict[str, str],
        endnotes: dict[str, str],
    ) -> list[Chapter]:
        """Walk through the document body and build Chapter/Section hierarchy."""
        chapters: list[Chapter] = []
        current_chapter_title = ""
        current_chapter_number: int | None = None
        current_section = Section(title="", level=2, paragraphs=[])
        current_sections: list[Section] = [current_section]
        chap_count = 0

        # Iterate over body elements (paragraphs and tables in order)
        for element in docx_doc.element.body:
            tag = element.tag.split("}")[-1] if "}" in element.tag else element.tag

            if tag == "tbl":
                # Table
                table = Table(element, docx_doc)
                text = self._table_to_text(table)
                if text.strip():
                    current_section.paragraphs.append(
                        Paragraph(
                            text=text,
                            metadata={"is_table": True},
                        )
                    )
            elif tag == "p":
                para = DocxParagraph(element, docx_doc)
                style_name = para.style.name if para.style else ""
                h_level = _heading_level_from_style(style_name)
                text = para.text.strip()

                if not text and h_level is None:
                    continue

                # Resolve footnote / endnote references
                note_ids = self._find_footnote_refs(para)
                note_texts: list[str] = []
                for nid in note_ids:
                    if nid in footnotes:
                        note_texts.append(footnotes[nid])
                    elif nid in endnotes:
                        note_texts.append(endnotes[nid])

                meta: dict[str, Any] = {}
                if note_texts:
                    meta["footnotes"] = note_texts
                    meta["is_footnote"] = False  # the para itself isn't a footnote

                if h_level == 1:
                    # Flush current chapter
                    if current_sections and (
                        any(s.paragraphs for s in current_sections)
                        or current_chapter_title
                    ):
                        chap_count += 1
                        chapters.append(
                            Chapter(
                                title=current_chapter_title,
                                number=current_chapter_number,
                                sections=current_sections,
                            )
                        )
                    current_chapter_title = text
                    current_chapter_number = chap_count + 1
                    current_section = Section(title="", level=2, paragraphs=[])
                    current_sections = [current_section]

                elif h_level == 2:
                    # New section within current chapter
                    current_section = Section(
                        title=text, level=2, paragraphs=[]
                    )
                    current_sections.append(current_section)

                elif h_level is not None and h_level >= 3:
                    # Sub-heading → paragraph with level metadata
                    meta["heading_level"] = h_level
                    current_section.paragraphs.append(
                        Paragraph(text=text, metadata=meta)
                    )

                else:
                    # Normal paragraph
                    current_section.paragraphs.append(
                        Paragraph(text=text, metadata=meta)
                    )

        # Flush last chapter
        chap_count += 1
        chapters.append(
            Chapter(
                title=current_chapter_title or "Untitled",
                number=current_chapter_number or chap_count,
                sections=current_sections,
            )
        )

        return chapters
